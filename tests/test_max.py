"""Tests run anywhere: no mic, GPU, Ollama or Windows needed."""
import sys
import threading
import time
from pathlib import Path

import numpy as np
import pytest

from max_assistant.agent import Agent, parse_yes_no, strip_thinking
from max_assistant.audio import FRAME, NoiseFloor, record_utterance
from max_assistant.config import Section, load_config
from max_assistant.llm import ChatReply, ToolCall
from max_assistant.main import Context
from max_assistant.tools import system as system_tools
from max_assistant.tools.registry import DECLINED, ToolRegistry
from max_assistant.tools.system import best_match
from max_assistant.tts import clean_for_speech


class FakeLLM:
    """Replays scripted replies and records what it was sent."""

    def __init__(self, replies):
        self.replies = list(replies)
        self.calls = []

    def chat(self, model, messages, tools=None):
        self.calls.append({"model": model, "messages": [dict(m) for m in messages],
                           "tools": [t["function"]["name"] for t in tools or []]})
        return self.replies.pop(0)

    def unload(self, model):
        self.unloaded = getattr(self, "unloaded", []) + [model]


def make_registry(tmp_path, llm=None):
    cfg = load_config()
    cfg["notes"] = {"file": str(tmp_path / "notes.md")}
    reg = ToolRegistry(context=Context(Section(cfg), llm))
    system_tools.register(reg)
    return reg


def make_agent(reg, llm, confirm=lambda p: False, events=None):
    return Agent(llm, reg, "fast", "planner", confirm=confirm,
                 on_event=(lambda k, d: events.append((k, d))) if events is not None else None)


# ---------- parsing ----------

@pytest.mark.parametrize("text,expected", [
    ("Yes", True), ("yeah go ahead", True), ("sure.", True), ("okay", True),
    ("no", False), ("yes... no wait, cancel", False), ("", False), (None, False),
    ("what?", False), ("don't do it", False),
])
def test_yes_no(text, expected):
    assert parse_yes_no(text) is expected


def test_strip_thinking():
    assert strip_thinking("<think>hmm\nplan</think> Hello") == "Hello"
    # Leaked reasoning with only the closing tag (seen after tool calls)
    assert strip_thinking("Okay, the user asked...\nSo: done.\n</think>\n\nSettings opened.") == "Settings opened."
    assert strip_thinking("Hi there <think>unfinished") == "Hi there"
    assert strip_thinking("Plain reply.") == "Plain reply."


def test_clean_for_speech():
    out = clean_for_speech("**Done!** See https://x.com\n- item 🎉")
    assert "*" not in out and "http" not in out and "🎉" not in out
    assert "Done!" in out and "a link" in out


def test_best_match():
    names = ["spotify", "google chrome", "visual studio code", "microsoft edge"]
    assert best_match("chrome", names) == "google chrome"
    assert best_match("spotfy", names) == "spotify"
    assert best_match("visual studio", names) == "visual studio code"
    assert best_match("photoshop", names) is None


# ---------- registry ----------

def test_registry_schemas_and_validation(tmp_path):
    reg = make_registry(tmp_path)
    names = {s["function"]["name"] for s in reg.schemas()}
    assert {"get_datetime", "open_app", "close_app", "take_note", "power", "free_gpu"} <= names
    assert reg.get("power").risky and reg.get("close_app").risky
    assert not reg.get("open_app").risky
    assert "missing argument" in reg.run("take_note", {})
    assert "unknown tool" in reg.run("hack_nasa", {})
    # hallucinated extra args are dropped, not passed through
    assert reg.run("get_datetime", {"timezone": "UTC"}).startswith("It's")


def test_notes_roundtrip(tmp_path):
    reg = make_registry(tmp_path)
    assert reg.run("read_notes", {}) == "You don't have any notes yet."
    reg.run("take_note", {"text": "buy milk"})
    reg.run("take_note", {"text": "call mom"})
    out = reg.run("read_notes", {"count": 5})
    assert "buy milk" in out and "call mom" in out


def test_open_website(tmp_path, monkeypatch):
    opened = []
    monkeypatch.setattr(system_tools.webbrowser, "open", opened.append)
    reg = make_registry(tmp_path)
    assert reg.run("open_website", {"target": "youtube.com"}) == "Opened youtube.com."
    assert reg.run("open_website", {"target": "weather in Tempe"}) == "Searching Google for weather in Tempe."
    assert reg.run("open_website", {"target": "cute dogs", "section": "images"}) ==         "Here are Google Images results for cute dogs."
    reg.run("open_website", {"target": "pizza", "section": "maps"})
    reg.run("open_website", {"target": "cats", "section": "bogus"})
    assert opened == [
        "https://youtube.com",
        "https://www.google.com/search?q=weather+in+Tempe",
        "https://www.google.com/search?q=cute+dogs&udm=2",
        "https://www.google.com/maps/search/pizza",
        "https://www.google.com/search?q=cats",            # unknown section -> normal search
    ]


def test_tool_exception_is_reported(tmp_path):
    reg = make_registry(tmp_path)

    @reg.tool("boom", params={})
    def explode():
        raise ValueError("kaboom")

    assert reg.run("explode", {}) == "Error running explode: kaboom"


# ---------- agent ----------

def test_plain_answer(tmp_path):
    llm = FakeLLM([ChatReply("Hello Manas.")])
    agent = make_agent(make_registry(tmp_path), llm)
    assert agent.handle("hi") == "Hello Manas."
    assert "think_harder" in llm.calls[0]["tools"]


def test_tool_call_then_answer(tmp_path):
    """Tools returning raw data (not direct) go back to the LLM to be phrased."""
    llm = FakeLLM([
        ChatReply("", [ToolCall("read_notes", {})]),
        ChatReply("You have no notes yet, Manas."),
    ])
    agent = make_agent(make_registry(tmp_path), llm)
    assert agent.handle("read my notes") == "You have no notes yet, Manas."
    tool_msg = llm.calls[1]["messages"][-1]
    assert tool_msg == {"role": "tool", "content": "You don't have any notes yet.", "tool_name": "read_notes"}


def test_direct_tool_skips_second_llm_call(tmp_path):
    llm = FakeLLM([ChatReply("", [ToolCall("take_note", {"text": "gym at 6"})])])
    agent = make_agent(make_registry(tmp_path), llm)
    assert agent.handle("note gym at 6") == "Noted."
    assert len(llm.calls) == 1
    assert "gym at 6" in (tmp_path / "notes.md").read_text()


def test_direct_tools_in_one_step_are_joined(tmp_path):
    reg = make_registry(tmp_path)
    reg.get("volume").func = lambda action, level=None: "Volume set to about 30 percent."
    llm = FakeLLM([ChatReply("", [ToolCall("take_note", {"text": "x"}),
                                  ToolCall("volume", {"action": "set", "level": 30})])])
    agent = make_agent(reg, llm)
    assert agent.handle("note x, volume 30") == "Noted. Volume set to about 30 percent."
    assert len(llm.calls) == 1


def test_compound_request_keeps_going_after_direct_tool(tmp_path):
    llm = FakeLLM([
        ChatReply("", [ToolCall("take_note", {"text": "gym"})]),
        ChatReply("Noted, and your notes are empty otherwise."),
    ])
    agent = make_agent(make_registry(tmp_path), llm)
    agent.handle("note gym and then read my notes")
    assert len(llm.calls) == 2   # "and"/"then" might need more steps, so the LLM continues


def test_direct_tool_error_goes_to_llm(tmp_path):
    llm = FakeLLM([
        ChatReply("", [ToolCall("volume", {"action": "set"})]),
        ChatReply("What level should I set it to?"),
    ])
    agent = make_agent(make_registry(tmp_path), llm)
    assert agent.handle("set the volume") == "What level should I set it to?"
    assert llm.calls[1]["messages"][-1]["content"].startswith("Error")


def test_system_prompt_is_stable_and_time_goes_in_user_message(tmp_path):
    llm = FakeLLM([ChatReply("Hi."), ChatReply("Sure.")])
    agent = make_agent(make_registry(tmp_path), llm)
    agent.handle("hello")
    agent.handle("thanks")
    first, second = llm.calls[0]["messages"], llm.calls[1]["messages"]
    assert first[0] == second[0]                       # identical system prompt -> reusable prefix
    assert first[1] == second[1]                       # history holds exactly what was sent
    assert first[-1]["content"].startswith("[") and first[-1]["content"].endswith("] hello")


def test_risky_tool_declined(tmp_path):
    ran = []
    reg = make_registry(tmp_path)
    reg.get("power").func = lambda action: ran.append(action)
    prompts = []
    llm = FakeLLM([ChatReply("", [ToolCall("power", {"action": "shutdown"})])])
    agent = make_agent(reg, llm, confirm=lambda p: prompts.append(p) or False)
    assert agent.handle("shut down") == "Okay, I won't."
    assert ran == []
    assert prompts == ["Do you really want me to shutdown the laptop?"]
    assert len(llm.calls) == 1   # a declined action needs no LLM phrasing


def test_risky_tool_approved(tmp_path):
    ran = []
    reg = make_registry(tmp_path)
    reg.get("power").func = lambda action: ran.append(action) or "Shutting down."
    llm = FakeLLM([
        ChatReply("", [ToolCall("power", {"action": "restart"})]),
        ChatReply("Restarting now."),
    ])
    events = []
    agent = make_agent(reg, llm, confirm=lambda p: True, events=events)
    agent.handle("restart")
    assert ran == ["restart"]
    kinds = [k for k, _ in events]
    assert kinds.index("confirm") < kinds.index("tool_call")


def test_escalation_switches_model(tmp_path):
    llm = FakeLLM([
        ChatReply("", [ToolCall("think_harder", {"reason": "planning"})]),
        ChatReply("Here's your study plan."),
    ])
    agent = make_agent(make_registry(tmp_path), llm)
    assert agent.handle("plan my week") == "Here's your study plan."
    assert [c["model"] for c in llm.calls] == ["fast", "planner"]
    assert "think_harder" not in llm.calls[1]["tools"]


def test_step_limit(tmp_path):
    llm = FakeLLM([ChatReply("", [ToolCall("read_notes", {})])] * 10)
    agent = make_agent(make_registry(tmp_path), llm)
    agent.max_steps = 3
    assert "stuck" in agent.handle("loop forever")
    assert len(llm.calls) == 3


def test_history_carries_over(tmp_path):
    llm = FakeLLM([ChatReply("It's 5 PM."), ChatReply("Sure.")])
    agent = make_agent(make_registry(tmp_path), llm)
    agent.handle("what time is it")
    agent.handle("and remind me later")
    roles = [m["role"] for m in llm.calls[1]["messages"]]
    assert roles == ["system", "user", "assistant", "user"]


def test_free_gpu_unloads_both(tmp_path):
    llm = FakeLLM([])
    reg = make_registry(tmp_path, llm)
    reg.run("free_gpu", {})
    cfg = load_config()
    assert set(llm.unloaded) == {cfg.llm.fast_model, cfg.llm.planner_model}


# ---------- audio ----------

class FakeMic:
    sample_rate = 16000

    def __init__(self, frames):
        self.frames = list(frames)

    def read(self, timeout=1.0):
        return self.frames.pop(0) if self.frames else None


def tone(amp):
    return (np.ones(FRAME) * amp).astype(np.int16)


def test_record_stops_after_silence():
    frames = [tone(10)] * 5 + [tone(5000)] * 10 + [tone(10)] * 20 + [tone(5000)] * 50
    audio = record_utterance(FakeMic(frames), NoiseFloor(100), silence_s=1.0, max_record_s=20)
    # lead-in (<=4) + 10 speech + ~13 silence frames; never reaches the later speech
    assert audio is not None
    assert len(audio) < 30 * FRAME
    assert audio.dtype == np.float32 and np.abs(audio).max() <= 1.0


def test_record_returns_none_when_silent():
    frames = [tone(10)] * 100
    assert record_utterance(FakeMic(frames), NoiseFloor(100), no_speech_timeout_s=2.0) is None


def test_wake_model_builtin():
    from max_assistant.wakeword import resolve_wake_model

    m = resolve_wake_model("hey_jarvis")
    assert (m.source, m.custom, m.phrase) == ("hey_jarvis", False, "Hey Jarvis")


def test_wake_model_custom_file(tmp_path):
    from max_assistant.wakeword import resolve_wake_model

    f = tmp_path / "hey_max.onnx"
    f.write_bytes(b"fake")
    m = resolve_wake_model(str(f), phrase="Hey Max")
    assert (m.source, m.custom, m.phrase) == (str(f), True, "Hey Max")
    assert resolve_wake_model(str(f)).phrase == "Hey Max"  # derived from the file name


def test_wake_model_missing_file_falls_back(tmp_path):
    from max_assistant.wakeword import resolve_wake_model

    m = resolve_wake_model(str(tmp_path / "nope.onnx"), phrase="Hey Max")
    assert (m.source, m.custom, m.phrase) == ("hey_jarvis", False, "Hey Jarvis")


def test_normalize_phrase():
    from max_assistant.wakeword import normalize_phrase

    assert normalize_phrase(" Hey, Max! ") == "HEY MAX"
    assert normalize_phrase("okay   max") == "OKAY MAX"


def test_keyword_line():
    from max_assistant.wakeword import keyword_line

    pieces = {"HEY MAX": ["▁HE", "Y", "▁MA", "X"]}
    line = keyword_line("Hey Max!", pieces.get, {"▁HE", "Y", "▁MA", "X"}, 1.0, 0.25)
    assert line == "▁HE Y ▁MA X :1.0 #0.25 @HEY_MAX"
    with pytest.raises(ValueError):
        keyword_line("Hey Max", pieces.get, {"▁HE"}, 1.0, 0.25)  # token missing from model
    with pytest.raises(ValueError):
        keyword_line("!!!", pieces.get, set(), 1.0, 0.25)


KWS_DIR = Path(__file__).resolve().parent.parent / "models" / "kws" / "gigaspeech-3.3M"


@pytest.mark.skipif(not (KWS_DIR / "tokens.txt").exists(), reason="keyword spotting model not downloaded")
@pytest.mark.parametrize("phrase,expected", [("lovely child", True), ("Hey Max", False)])
def test_sherpa_detector_on_real_audio(phrase, expected):
    """End to end on the model's bundled clip, which contains '...a lovely child...'."""
    pytest.importorskip("sherpa_onnx")
    import wave

    from max_assistant.wakeword import SherpaKeywordDetector

    det = SherpaKeywordDetector(phrase, str(KWS_DIR))
    with wave.open(str(KWS_DIR / "test_wavs" / "1.wav")) as w:
        assert w.getframerate() == 16000
        audio = np.frombuffer(w.readframes(w.getnframes()), dtype=np.int16)
    audio = np.concatenate([audio, np.zeros(16000, np.int16)])
    heard = any(det.process(audio[i:i + FRAME]) for i in range(0, len(audio), FRAME))
    assert heard is expected


def test_auto_gain_boosts_quiet_but_not_loud():
    from max_assistant.wakeword import AutoGain

    agc = AutoGain(max_gain=8.0, target_peak=16000, window_frames=3)
    quiet = np.full(FRAME, 500, np.int16)
    assert agc(quiet).max() == pytest.approx(4000)          # capped at 8x
    loud = np.full(FRAME, 20000, np.int16)
    assert agc(loud).max() == pytest.approx(20000)          # never turned down
    assert agc(quiet).max() == pytest.approx(500)           # loud word still in window: no pumping
    for _ in range(3):
        out = agc(quiet)
    assert out.max() == pytest.approx(4000)                 # window passed: boost returns


class FakeStream:
    def __init__(self, sid):
        self.id, self.samples, self.pending, self.fired = sid, 0, False, False

    def accept_waveform(self, sr, x):
        self.samples += len(x)
        self.pending = True


class FakeSpotter:
    """Stands in for sherpa_onnx.KeywordSpotter: only stream `fire_stream` fires, once it has
    seen `fire_after` samples."""

    def __init__(self, fire_stream=2, fire_after=FRAME):
        self.fire_stream, self.fire_after, self.created = fire_stream, fire_after, []

    def create_stream(self):
        self.created.append(FakeStream(len(self.created)))
        return self.created[-1]

    def is_ready(self, s):
        return s.pending

    def decode_stream(self, s):
        s.pending = False

    def get_result(self, s):
        if s.id == self.fire_stream and not s.fired and s.samples >= self.fire_after:
            s.fired = True
            return "HEY_MAX"
        return ""

    def reset_stream(self, s):
        pass


def test_sherpa_streams_start_staggered_and_any_can_fire():
    from max_assistant.wakeword import SherpaKeywordDetector, WakeDetector

    det = SherpaKeywordDetector.__new__(SherpaKeywordDetector)   # skip model loading
    WakeDetector.__init__(det, cooldown_s=1.0)
    det.spotter, det.agc, det.num_streams = FakeSpotter(), None, 4
    det._stagger = int(16000 * det.CHUNK_S / det.num_streams)
    det.reset()

    started = []                      # (frame, streams alive) before any detection
    fired = []
    for i in range(12):
        if det.process(np.zeros(FRAME, np.int16)):
            fired.append(i)
            assert det.streams == [None] * 4   # firing restarts the staggered set
            break
        started.append(sum(s is not None for s in det.streams))
    # 0.64 s chunk / 4 streams = 2560 samples = 2 frames: a new stream every 2 frames
    assert det._stagger == 2560
    assert started == [1, 1, 2, 2]
    assert fired == [4]               # stream #2 starts at frame 4 and fires on its first frame


def test_llm_forces_gpu_layers_and_falls_back_on_load_error(monkeypatch):
    from max_assistant import llm as llm_mod

    sent = []

    class Resp:
        def __init__(self, code, body):
            self.status_code, self._body, self.text = code, body, str(body)

        def json(self):
            return self._body

    def fake_post(url, json, timeout):
        sent.append(json["options"].copy())
        if "num_gpu" in json["options"]:
            return Resp(500, {"error": "cudaMalloc failed: out of memory"})
        return Resp(200, {"message": {"content": "hi"}})

    monkeypatch.setattr(llm_mod.requests, "post", fake_post)
    client = llm_mod.OllamaClient("http://x", num_ctx=4096, num_predict=256, gpu_layers={"fast": 99})
    assert client.chat("fast", [{"role": "user", "content": "hey"}]).content == "hi"
    assert sent[0]["num_gpu"] == 99 and sent[0]["num_ctx"] == 4096 and sent[0]["num_predict"] == 256
    assert "num_gpu" not in sent[1]            # retried with Ollama's own split
    client.chat("fast", [{"role": "user", "content": "again"}])
    assert "num_gpu" not in sent[2]            # and stays off for the session
    client.chat("planner", [])
    assert "num_gpu" not in sent[3]            # never forced for other models


# ---------- web ----------

def test_select_passages_prefers_query_terms():
    from max_assistant.web import select_passages

    text = ("Paris is the capital of France. It has many museums. " * 3 +
            "The Eiffel Tower is 330 metres tall including antennas. " +
            "Bread is popular. " * 10)
    picked = select_passages(text, "how tall is the Eiffel Tower", n=1, size=120)
    assert len(picked) == 1 and "330 metres" in picked[0]


def test_select_passages_recency_finds_last_table_rows():
    from max_assistant.web import select_passages

    rows = " ".join(f"| {1967 + i} | Team{i} won the Super Bowl |." for i in range(60))
    picked = select_passages(rows, "most recent Super Bowl winner", n=1, size=200, this_year=2026)
    assert "2026" in picked[0] and "Team59" in picked[0]
    old = select_passages(rows, "first Super Bowl winner", n=1, size=200, this_year=2026)
    assert "1967" in old[0]                      # no recency intent: earliest match wins


def test_web_search_falls_back_to_duckduckgo(monkeypatch):
    import requests

    from max_assistant import web

    def boom(*a, **k):
        raise requests.ConnectionError("docker is off")

    monkeypatch.setattr(web.requests, "get", boom)
    monkeypatch.setattr(web.WebSearch, "_duckduckgo",
                        lambda self, q, n: [web.SearchResult("DDG hit", "https://example.com/a", "snippet")])
    ws = web.WebSearch("http://127.0.0.1:8888")
    hits = ws.search("anything")
    assert [h.title for h in hits] == ["DDG hit"] and ws.last_backend == "duckduckgo"


def test_web_search_tool_includes_passages_and_guard(tmp_path, monkeypatch):
    from max_assistant import web
    from max_assistant.tools import web as web_tools

    monkeypatch.setattr(web.WebSearch, "search", lambda self, q, n=None: [
        web.SearchResult("Super Bowl Winners", "https://www.espn.com/sb", "Every champion since 1967.")])
    monkeypatch.setattr(web, "fetch_text", lambda url, timeout=5.0: "| 2025 | Philadelphia won |. | 2026 | Seattle won |.")
    cfg = load_config()
    reg = ToolRegistry(context=Context(cfg, FakeLLM([])))
    web_tools.register(reg)
    out = reg.run("web_search", {"query": "latest Super Bowl winner"})
    assert "1. Super Bowl Winners (espn.com)" in out
    assert "Seattle won" in out                  # passage read from the page
    assert "don't guess" in out                  # anchored to the results


# ---------- browser (fake page, no Chrome needed) ----------

def el(i, name, role="link", **kw):
    from max_assistant.browser import Element

    base = dict(tag="a" if role == "link" else "button", type="", in_view=True, top=100 * i, heading=False,
                search=False, left=0, image=False, main=True)
    base.update(kw)
    return Element(i, name, base.pop("tag"), base.pop("type"), role, **base)


def test_pick_exact_ordinal_number_and_ambiguous():
    from max_assistant.browser import pick

    page = [el(1, "Images", role="tab", main=False), el(2, "Videos", role="tab", main=False),
            el(3, "Cute puppies playing in the snow", heading=True), el(4, "Ten tiny dogs you will love", heading=True),
            el(5, "Privacy", main=False)]
    assert pick(page, "Images")[0].id == 1
    assert pick(page, "first result")[0].id == 3            # nav tabs aren't results
    assert pick(page, "second result")[0].id == 4
    assert pick(page, "#5")[0].id == 5
    chosen, candidates = pick(page, "zzz qqq")
    assert chosen is None and candidates                    # nothing clear: let the LLM choose


def test_pick_thumbnail_becomes_title_link():
    from max_assistant.browser import pick

    page = [el(1, "Sponsored thing you should buy today", top=50),
            el(2, "11:53:45 Now playing", top=400, image=True),
            el(3, "lofi hip hop radio - beats to relax/study to", top=402),
            el(4, "6:10:58", top=700, image=True),
            el(5, "Best of lofi hip hop 2021 - 6 hour mix", top=703)]
    assert pick(page, "first video")[0].id == 3             # ad has no thumbnail; title beats "11:53:45"
    assert pick(page, "second video")[0].id == 5


class FakeSession:
    running = True

    def __init__(self, elements):
        self.elements, self.clicked, self.typed = elements, [], []

    def scan(self):
        return self.elements

    def settle(self, *a):
        pass

    def click(self, e):
        self.clicked.append(e.id)

    def type_into(self, e, text, submit):
        self.typed.append((e.id if e else None, text, submit))


def browser_registry(confirm_answer=False):
    from max_assistant.tools import browser as browser_tools

    cfg = load_config()
    ctx = Context(cfg, FakeLLM([]))
    asked = []
    ctx.confirm = lambda p: asked.append(p) or confirm_answer
    reg = ToolRegistry(context=ctx)
    browser_tools.register(reg)              # picks up the FakeSession patched in by patch_session
    return reg, ctx, asked


def patch_session(monkeypatch, fake):
    from max_assistant.tools import browser as browser_tools

    monkeypatch.setattr(browser_tools, "BrowserSession", lambda *a, **k: fake)


def test_browser_click_risky_needs_yes(tmp_path, monkeypatch):
    fake = FakeSession([el(1, "Subscribe to Lofi Girl.", role="button"), el(2, "Share", role="button")])
    patch_session(monkeypatch, fake)
    reg, ctx, asked = browser_registry(confirm_answer=False)
    assert reg.run("browser_click", {"target": "Share"}) == "Clicked Share."
    assert reg.run("browser_click", {"target": "Subscribe"}) == DECLINED
    assert asked == ["This will click 'Subscribe to Lofi Girl'. Should I?"]
    assert fake.clicked == [2]


def test_browser_type_search_submits_and_refuses_passwords(tmp_path, monkeypatch):
    fake = FakeSession([el(1, "Search", role="searchbox", tag="input", search=True),
                        el(2, "Password", role="textbox", tag="input", type="password"),
                        el(3, "Comment", role="textbox", tag="textarea")])
    patch_session(monkeypatch, fake)
    reg, ctx, asked = browser_registry()
    assert reg.run("browser_type", {"text": "lofi music"}) == "Typed it and submitted."
    assert fake.typed == [(1, "lofi music", True)]           # search box: Enter by default
    assert reg.run("browser_type", {"text": "hunter2", "field": "password"}).startswith("Error: I don't type passwords")
    assert reg.run("browser_type", {"text": "great video", "field": "comment", "submit": True}) == DECLINED
    assert len(fake.typed) == 1 and len(asked) == 1          # posting a comment needed a yes


def test_ambiguous_click_goes_back_to_llm(tmp_path, monkeypatch):
    fake = FakeSession([el(1, "Open settings panel", role="button"), el(2, "Open settings menu", role="button")])
    patch_session(monkeypatch, fake)
    reg, ctx, asked = browser_registry()
    llm = FakeLLM([ChatReply("", [ToolCall("browser_click", {"target": "zzqq"})]),
                   ChatReply("Which one: the settings panel or the settings menu?")])
    agent = make_agent(reg, llm)
    assert agent.handle("click that thing") == "Which one: the settings panel or the settings menu?"
    assert len(llm.calls) == 2                               # list went to the LLM, not read aloud


def test_pick_counts_videos_on_a_video_page():
    from max_assistant.browser import pick

    page = [el(1, "The Beatles", image=True, href="/channel/UCc4K7"),
            el(2, "Hey Jude (Remastered 2015)", image=True, href="/watch?v=A_MjCqQoLLA"),
            el(3, "Let It Be (Remastered 2009)", image=True, href="/watch?v=QDYfEBY9NM4"),
            el(4, "Come Together", image=True, href="/watch?v=45cYwDMibGo")]
    assert pick(page, "second link")[0].id == 3            # channel card isn't counted
    assert pick(page, "first video")[0].id == 2


def test_open_website_site_search(tmp_path, monkeypatch):
    opened = []
    monkeypatch.setattr(system_tools.webbrowser, "open", opened.append)
    reg = make_registry(tmp_path)
    assert reg.run("open_website", {"target": "Beatles", "site": "youtube"}) == "Searching YouTube for Beatles."
    reg.run("open_website", {"target": "usb c hub", "site": "amazon.com"})
    assert opened == ["https://www.youtube.com/results?search_query=Beatles", "https://www.amazon.com/s?k=usb+c+hub"]


def test_background_writer_never_blocks_the_caller():
    import threading
    import time

    from max_assistant.winutil import BackgroundWriter

    class StalledConsole:
        """Like a console in 'Select' mode: writes block until released."""
        def __init__(self):
            self.release, self.written = threading.Event(), []

        def write(self, text):
            self.release.wait()
            self.written.append(text)

        def flush(self):
            pass

    console = StalledConsole()
    out = BackgroundWriter(console, max_items=10)
    start = time.perf_counter()
    for i in range(50):
        print(f"line {i}", file=out)       # would hang forever if written directly
    assert time.perf_counter() - start < 0.5
    console.release.set()
    for _ in range(50):
        if console.written:
            break
        time.sleep(0.02)
    assert console.written and console.written[0] == "line 0"


# ---------- listening ----------

def test_looks_complete_and_strip_wake_word():
    from max_assistant.listen import looks_complete, strip_wake_word

    assert looks_complete("Open YouTube.") and looks_complete("What time is it?")
    assert not looks_complete("Can you search for...") and not looks_complete("Search for")
    assert not looks_complete("open the") and not looks_complete("")
    assert strip_wake_word("Max, open YouTube.") == "open YouTube."
    assert strip_wake_word("Hey Max what time is it") == "what time is it"
    assert strip_wake_word("Maxwell's equations") == "Maxwell's equations"


def speech(n, amp=5000):
    return [tone(amp)] * n


def test_listen_answers_early_when_the_request_is_complete():
    from max_assistant.listen import listen_for_command

    calls = []

    def transcribe(audio):
        calls.append(len(audio))
        return "Open YouTube."

    frames = speech(10) + [tone(10)] * 40            # 0.8 s speech, then 3.2 s of silence
    mic = FakeMic(frames)
    assert listen_for_command(mic, NoiseFloor(100), transcribe, silence_s=1.0, early_s=0.4) == "Open YouTube."
    used = 50 - len(mic.frames)
    assert used < 10 + 12 + 4                        # returned at ~0.4-0.5 s of silence, not 1.0 s


def test_listen_waits_when_the_sentence_trails_off():
    from max_assistant.listen import listen_for_command

    replies = iter(["Can you search for...", "Can you search for arctic monkeys?", "Can you search for arctic monkeys?"])
    frames = speech(10) + [tone(10)] * 15 + speech(8) + [tone(10)] * 40   # 1.2 s pause mid-sentence
    mic = FakeMic(frames)
    text = listen_for_command(mic, NoiseFloor(100), lambda a: next(replies), silence_s=1.0, early_s=0.4,
                              max_silence_s=1.8)
    assert text == "Can you search for arctic monkeys?"   # didn't cut off at the pause


def test_listen_keeps_words_said_right_after_wake():
    from max_assistant.listen import listen_for_command

    got = []
    frames = speech(6) + [tone(10)] * 30             # talking immediately, no pause for the chime
    listen_for_command(FakeMic(frames), NoiseFloor(100), lambda a: got.append(a) or "Open Spotify.")
    assert got and len(got[0]) >= 6 * FRAME          # all the speech frames were transcribed


def test_stt_routes_clear_speech_to_fast_engine_and_quiet_to_whisper():
    from max_assistant.stt import SpeechToText, snr_db

    rng = np.random.default_rng(0)
    noise = lambda n, db: rng.normal(0, 10 ** (db / 20), n)
    tone_ = lambda n, db: np.sin(np.arange(n) / 5) * 10 ** (db / 20)
    clear = np.concatenate([noise(8000, -60), tone_(16000, -6) + noise(16000, -60), noise(8000, -60)]).astype(np.float32)
    murky = np.concatenate([noise(8000, -40), tone_(16000, -38) + noise(16000, -40), noise(8000, -40)]).astype(np.float32)
    assert snr_db(clear) > 40 and snr_db(murky) < 12

    stt = SpeechToText.__new__(SpeechToText)       # skip loading real models
    used = []
    stt.fast = type("F", (), {"transcribe": lambda self, a: used.append("fast") or "Open YouTube."})()
    stt.fast_min_snr_db = 12.0
    stt.whisper = lambda a: used.append("whisper") or "open youtube"
    assert stt.transcribe(clear) == "Open YouTube." and used == ["fast"]
    assert stt.transcribe(murky) == "open youtube" and used == ["fast", "whisper"]
    stt.fast = type("F", (), {"transcribe": lambda self, a: ""})()   # fast engine heard nothing
    assert stt.transcribe(clear) == "open youtube"                   # -> Whisper double-checks


def test_listen_retranscribes_everything_when_it_still_sounds_unfinished():
    from max_assistant.listen import listen_for_command

    seen = []

    def transcribe(audio):
        seen.append(len(audio))
        return "search for arctic monkeys on..." if len(seen) == 1 else "search for arctic monkeys on YouTube"

    frames = speech(10) + [tone(10)] * 40
    text = listen_for_command(FakeMic(frames), NoiseFloor(100), transcribe, silence_s=1.0, early_s=0.4, max_silence_s=1.8)
    assert text == "search for arctic monkeys on YouTube"
    assert len(seen) == 2 and seen[1] > seen[0]      # second pass covered the quiet tail too


# ---------- follow-up conversation ----------

def run_converse(first, replies, follow_ups, follow_up=True):
    from max_assistant.main import converse

    replies, follow_ups, said, idle = iter(replies), iter(follow_ups), [], []

    def respond(text):
        said.append(text)
        return next(replies) if text is not None else None

    n = converse(first, respond, lambda: next(follow_ups), follow_up=follow_up, on_idle=lambda: idle.append(1))
    return n, said, idle


def test_follow_up_after_a_question_without_wake_word():
    n, said, idle = run_converse("search for something",
                                 ["What should I search for?", "Searching Google for lofi music."],
                                 ["lofi music"])
    assert said == ["search for something", "lofi music"] and n == 2


def test_no_follow_up_after_a_statement():
    n, said, _ = run_converse("open youtube", ["Opened youtube.com."], [])   # follow_ups would raise if used
    assert n == 1


def test_follow_up_ends_on_silence_or_phantom_text():
    n, said, idle = run_converse("play something", ["Which song?"], [""])
    assert n == 1 and idle == [1]                      # silence -> back to waiting for "Hey Max"
    n, said, idle = run_converse("play something", ["Which song?"], ["Thank you."])
    assert n == 1 and idle == [1]                      # classic phantom transcript ignored


def test_follow_ups_chain_and_can_be_turned_off():
    n, said, _ = run_converse("set a reminder", ["For when?", "What should it say?", "Noted."],
                              ["tomorrow", "call mom"])
    assert said == ["set a reminder", "tomorrow", "call mom"] and n == 3
    n, said, _ = run_converse("search", ["What for?"], [], follow_up=False)
    assert n == 1


def test_never_mind_in_a_follow_up():
    n, said, _ = run_converse("search", ["What should I search for?"], ["never mind"])
    assert said == ["search", None] and n == 1        # None = the "Okay." path


def test_snr_not_fooled_by_noise_gated_mic():
    from max_assistant.stt import snr_db

    gated = np.concatenate([np.zeros(8000), np.sin(np.arange(16000) / 5) * 0.3, np.zeros(8000)]).astype(np.float32)
    assert snr_db(gated) < 70            # digital-zero "silence" no longer means a perfect recording


# ---------- memory & reminders ----------

class WordEmbedder:
    """Deterministic stand-in for the real embedder: hashed bag of words, normalised."""

    def embed(self, texts):
        import re as _re
        out = np.zeros((len(texts), 256), dtype=np.float32)
        for i, t in enumerate(texts):
            for w in _re.findall(r"[a-z0-9]+", t.lower()):
                if w not in {"my", "is", "the", "a", "on", "what", "when", "s"}:
                    out[i, hash(w) % 256] += 1
        norms = np.linalg.norm(out, axis=1, keepdims=True)
        return out / np.where(norms == 0, 1, norms)


def memory_store():
    from max_assistant.memory import MemoryStore
    return MemoryStore(":memory:", WordEmbedder())


def test_memory_facts_search_dedupe_delete():
    m = memory_store()
    fid, old = m.add_fact("My CSE 572 exam is on Friday October 2.")
    assert old is None
    m.add_fact("My sister's name is Priya")
    hits = m.search_facts("when is my CSE 572 exam", k=1)
    assert hits[0].text == "My CSE 572 exam is on Friday October 2"
    _, old = m.add_fact("My CSE 572 exam is on Friday October 2")        # same fact again
    assert old is not None and len(m.facts()) == 2
    assert m.delete_fact(fid) and len(m.facts()) == 1
    assert not m.search_facts("CSE 572 exam", min_score=0.5)


def test_memory_turns_and_actions():
    m = memory_store()
    m.log_turn("search for arctic monkeys", "Searching YouTube for arctic monkeys.", ["open_website"])
    m.log_turn("set volume to 30", "Volume set to about 30 percent.", ["volume"])
    assert m.search_turns("arctic monkeys", k=1)[0].text.startswith("you: search for arctic monkeys")
    assert [t["tools"] for t in m.recent_turns()] == [["open_website"], ["volume"]]
    m.log_action("tool", {"tool": "volume", "result": "ok"})
    assert m.actions()[-1]["detail"]["tool"] == "volume"


def test_context_for_only_attaches_relevant_facts():
    m = memory_store()
    m.add_fact("My CSE 572 exam is on Friday October 2")
    assert "exam" in (m.context_for("CSE 572 exam Friday October") or "")
    assert m.context_for("open youtube") is None


def test_parse_when_phrases():
    import datetime as dt
    from max_assistant.reminders import parse_when, spoken_time

    now = dt.datetime(2026, 9, 30, 0, 45)          # Wednesday, 12:45 AM
    cases = {"Friday at 9am": (2026, 10, 2, 9, 0), "friday": (2026, 10, 2, 9, 0),
             "tomorrow morning": (2026, 10, 1, 9, 0), "tonight at 9": (2026, 9, 30, 21, 0),
             "next Monday 8am": (2026, 10, 5, 8, 0), "in 2 hours": (2026, 9, 30, 2, 45),
             "Thursday evening": (2026, 10, 1, 18, 0), "2026-10-02T09:00": (2026, 10, 2, 9, 0)}
    for phrase, expect in cases.items():
        assert parse_when(phrase, now) == dt.datetime(*expect), phrase
    assert parse_when("gibberish zzz", now) is None
    assert spoken_time(dt.datetime(2026, 10, 1, 18, 0), now) == "tomorrow at 6 PM"
    assert spoken_time(dt.datetime(2026, 10, 2, 9, 30), now) == "Friday, October 2 at 9:30 AM"


def test_reminder_scheduler_fires_due_and_flags_missed():
    import datetime as dt
    from max_assistant.reminders import ReminderScheduler, Reminders

    r = Reminders(memory_store())
    now = dt.datetime.now()
    r.add("stretch", now - dt.timedelta(seconds=30))       # just due
    r.add("exam", now - dt.timedelta(hours=3))             # came due while Max was off
    r.add("later", now + dt.timedelta(hours=1))
    fired = []
    ReminderScheduler(r, lambda rem, missed: fired.append((rem.text, missed))).check(now)
    assert sorted(fired) == [("exam", True), ("stretch", False)]
    assert [x.text for x in r.pending()] == ["later"]      # fired ones aren't repeated


def memory_registry(confirm_answer=True):
    from max_assistant.reminders import Reminders
    from max_assistant.tools import memory as memory_tools

    ctx = Context(load_config(), FakeLLM([]))
    ctx.memory = memory_store()
    ctx.reminders = Reminders(ctx.memory)
    ctx.confirm = lambda p: confirm_answer
    reg = ToolRegistry(context=ctx)
    memory_tools.register(reg)
    return reg, ctx


def test_memory_tools_end_to_end():
    reg, ctx = memory_registry()
    assert reg.run("remember", {"fact": "My CSE 572 exam is on Friday October 2"}) == "Got it, I'll remember that."
    out = reg.run("recall", {"query": "when is my CSE 572 exam"})
    assert "Friday October 2" in out
    assert reg.run("forget", {"what": "CSE 572 exam"}) == "Forgotten."
    assert reg.run("recall", {"query": "CSE 572 exam"}) == "Nothing saved about that."


def test_reminder_tools():
    reg, ctx = memory_registry()
    out = reg.run("set_reminder", {"text": "CSE 572 exam", "when": "tomorrow at 9am"})
    assert out.startswith("Okay, I'll remind you tomorrow at 9 AM")
    assert reg.run("set_reminder", {"text": "x", "when": "blorp"}).startswith("Error")
    reg.run("set_reminder", {"text": "call mom", "when": "in 2 hours"})
    assert reg.run("list_reminders", {}).startswith("You have 2 reminders")
    assert reg.run("cancel_reminder", {"which": "the exam"}) == "Cancelled the reminder: CSE 572 exam."
    assert [r.text for r in ctx.reminders.pending()] == ["call mom"]


def test_agent_attaches_memories_and_logs_turns(tmp_path):
    logged = []
    llm = FakeLLM([ChatReply("", [ToolCall("take_note", {"text": "x"})])])
    agent = Agent(llm, make_registry(tmp_path), "fast", "planner",
                  recall=lambda t: "Manas's exam is on Friday" if "exam" in t else None,
                  on_turn=lambda u, a, tools: logged.append((u, a, tools)))
    agent.handle("note about my exam")
    assert "(You remember: Manas's exam is on Friday)" in llm.calls[0]["messages"][-1]["content"]
    assert logged == [("note about my exam", "Noted.", ["take_note"])]


# ---------- dashboard API ----------

LOCAL = ("127.0.0.1", 50000)       # requests from the laptop itself (no token needed)

def api_client(run_command=lambda text, source="dashboard": f"echo: {text}", token="", client=LOCAL):
    from fastapi.testclient import TestClient

    from max_assistant.events import ApprovalBroker, EventBus
    from max_assistant.reminders import Reminders
    from max_assistant.server import Runtime, create_app

    ctx = Context(load_config(), FakeLLM([]))
    ctx.memory = memory_store()
    ctx.reminders = Reminders(ctx.memory)
    from max_assistant.agenda import Events

    ctx.events = Events(ctx.memory)
    bus = EventBus()
    approvals = ApprovalBroker(bus)
    rt = Runtime(ctx, bus, approvals, run_command, {"name": "Max"}, token=token)
    return TestClient(create_app(rt), client=client), rt


def test_api_facts_reminders_and_commands():
    client, rt = api_client()
    assert client.post("/api/facts", json={"text": "My sister's name is Priya"}).json()["updated"] is False
    facts = client.get("/api/facts").json()
    assert [f["text"] for f in facts] == ["My sister's name is Priya"]
    assert client.delete(f"/api/facts/{facts[0]['id']}").json() == {"ok": True}
    r = client.post("/api/reminders", json={"text": "exam", "when": "tomorrow at 9am"}).json()
    assert r["spoken"] == "tomorrow at 9 AM"
    assert [x["text"] for x in client.get("/api/reminders").json()] == ["exam"]
    assert client.post("/api/reminders", json={"text": "x", "when": "blorp"}).status_code == 400
    assert client.post("/api/command", json={"text": "open youtube"}).json() == {"reply": "echo: open youtube"}
    assert [e["kind"] for e in rt.bus.recent][-1] == "reminder"


def test_api_approvals_race_and_websocket():
    client, rt = api_client()
    a = rt.approvals.open("Shut down the laptop?", "power")
    assert client.get("/api/state").json()["approvals"] == [{"id": a.id, "prompt": "Shut down the laptop?", "tool": "power"}]
    assert client.post(f"/api/approvals/{a.id}", json={"approved": True}).json() == {"ok": True}
    assert a.done.is_set() and a.approved is True
    rt.approvals.close(a, True, "dashboard")
    assert client.post(f"/api/approvals/{a.id}", json={"approved": False}).status_code == 404
    with client.websocket_connect("/api/ws") as ws:
        hello = ws.receive_json()
        assert hello["kind"] == "hello" and any(e["kind"] == "approval_result" for e in hello["data"]["recent"])
        rt.bus.publish("stage", {"stage": "thinking"})
        assert ws.receive_json()["data"] == {"stage": "thinking"}


def test_forget_requests_skip_the_model_and_still_confirm():
    from max_assistant.agent import FORGET

    assert FORGET.match("could you forget my exam").group(1) == "my exam"
    assert FORGET.match("can you forget what I told you about lofi").group(1) == "lofi"
    assert FORGET.match("I forgot my keys") is None and FORGET.match("don't forget to call mom") is None
    reg, ctx = memory_registry(confirm_answer=False)
    ctx.memory.add_fact("Manas prefers lofi music when studying")
    llm = FakeLLM([])                                   # would raise if the model were asked
    agent = Agent(llm, reg, "fast", "planner")
    assert agent.handle("forget the lofi music when studying") == "Okay, I won't."
    assert len(ctx.memory.facts()) == 1 and llm.calls == []


def test_false_claim_gets_one_nudge():
    reg, ctx = memory_registry()
    llm = FakeLLM([ChatReply("Okay, I'll remind you tomorrow."),               # claims, no tool
                   ChatReply("", [ToolCall("set_reminder", {"text": "gym", "when": "tomorrow at 7am"})])])
    agent = Agent(llm, reg, "fast", "planner")
    assert agent.handle("remind me about gym tomorrow at 7").startswith("Okay, I'll remind you tomorrow at 7 AM")
    assert "System check" in llm.calls[1]["messages"][-1]["content"]
    assert [r.text for r in ctx.reminders.pending()] == ["gym"]


# ---------- Canvas & daily digest ----------

ICS = b"""BEGIN:VCALENDAR
VERSION:2.0
PRODID:test
BEGIN:VEVENT
UID:event-assignment-1
DTSTART;VALUE=DATE:20261001
SUMMARY:Lab 4 [2026FallC-T-CSE572-68549]
END:VEVENT
BEGIN:VEVENT
UID:event-assignment-2
DTSTART;VALUE=DATE:20261004
SUMMARY:Assignment1 [2026FallC-T-CSE579-73246]
END:VEVENT
BEGIN:VEVENT
UID:event-calendar-event-10
DTSTART:20261001T120000Z
DTEND:20261001T131500Z
SUMMARY:CSE 573: Semantic Web Mining (2026 Fall C) [2026FallC-T-CSE573-82398]
DESCRIPTION:[Click here to join Zoom Meeting] (https://asu.zoom.us/j/123)
END:VEVENT
BEGIN:VEVENT
UID:event-calendar-event-11
DTSTART:20261001T120000Z
DTEND:20261001T131500Z
SUMMARY:CSE 573: Semantic Web Mining (2024 Spring) [2026FallC-T-CSE573-82398]
END:VEVENT
END:VCALENDAR
"""


def test_canvas_feed_parsing():
    from max_assistant.canvas import describe_due, parse_feed

    items = parse_feed(ICS)
    assignments = [i for i in items if i.kind == "assignment"]
    classes = [i for i in items if i.kind == "class"]
    assert [(a.title, a.course, a.all_day) for a in assignments] == [("Lab 4", "CSE 572", True), ("Assignment1", "CSE 579", True)]
    assert len(classes) == 1                                  # other section's copy dropped
    assert classes[0].course == "CSE 573" and classes[0].link == "https://asu.zoom.us/j/123"
    assert describe_due(assignments[:1], "today") == "Due today: Lab 4 for CSE 572, by end of day."
    assert describe_due([], "tomorrow") == "Nothing due tomorrow on Canvas."


class FakeFeed:
    def __init__(self, items):
        self._items = items

    def between(self, start, end, kind=None):
        return [i for i in self._items if start <= i.start < end and (kind is None or i.kind == kind)]


def digest_ctx(tmp_items=None):
    import datetime as dt
    from max_assistant.canvas import parse_feed
    from max_assistant.reminders import Reminders

    ctx = Context(load_config(), FakeLLM([]))
    ctx.memory = memory_store()
    ctx.reminders = Reminders(ctx.memory)
    ctx.canvas = FakeFeed(parse_feed(ICS) if tmp_items is None else tmp_items)
    ctx.reminders.add("stand-up", dt.datetime(2026, 10, 1, 10, 0))
    return ctx


def test_digest_content_and_spoken_summary():
    import datetime as dt
    from max_assistant.digest import gather, render, spoken

    plan = gather(digest_ctx(), dt.date(2026, 10, 1))
    assert [i.title for i in plan.due_today] == ["Lab 4"] and [i.title for i in plan.due_soon] == ["Assignment1"]
    assert len(plan.classes) == 1 and [r.text for r in plan.reminders] == ["stand-up"]
    subject, text, body = render(plan, "Focus on Lab 4 first.")
    assert subject == "Your day, Thu Oct 1: 1 due today"
    assert "Lab 4 (CSE 572) by end of day" in text and "Zoom: https://asu.zoom.us/j/123" in text
    assert "Focus on Lab 4 first." in body and "<li" in body
    s = spoken(plan)
    assert s.startswith("Due today: Lab 4 for CSE 572.") and "Coming up: Assignment1" in s


def test_digest_schedule_catch_up_and_once_a_day():
    import datetime as dt
    from max_assistant.digest import DigestScheduler

    ctx = digest_ctx()
    sched = DigestScheduler(ctx, "07:00", weekends=False, catch_up_until="18:00")
    thu = dt.datetime(2026, 10, 1, 6, 59)
    assert not sched.due(thu)                                       # too early
    assert sched.due(thu.replace(hour=9))                           # late start: catch up
    assert not sched.due(thu.replace(hour=19))                      # too late to be useful
    assert not sched.due(dt.datetime(2026, 10, 3, 8, 0))            # Saturday, weekends off
    ctx.memory.set("digest_last_sent", "2026-10-01")
    assert not sched.due(thu.replace(hour=9))                       # already sent today


def test_digest_refuses_placeholder_email_settings():
    from max_assistant.digest import send_email

    with pytest.raises(RuntimeError, match="isn't set up"):
        send_email({"smtp_user": "your.address@gmail.com", "app_password": "xxxx xxxx xxxx xxxx"}, "s", "t", "<p>h</p>")


def test_api_remote_control_and_today():
    from fastapi.testclient import TestClient

    from max_assistant.events import ApprovalBroker, Controls, EventBus
    from max_assistant.server import Runtime, create_app

    ctx = digest_ctx()
    bus = EventBus()
    controls = Controls()
    client = TestClient(create_app(Runtime(ctx, bus, ApprovalBroker(bus), lambda t: t, {}, controls)), client=LOCAL)
    controls.paused.set()
    assert client.post("/api/control/listen").json() == {"ok": True, "paused": False}
    assert controls.listen_now.is_set() and not controls.paused.is_set()     # push-to-talk also unpauses
    assert client.post("/api/control/pause").json()["paused"] is True
    client.post("/api/control/quit")
    assert controls.quit.is_set()
    assert client.post("/api/control/explode").status_code == 404
    today = client.get("/api/today").json()
    assert set(today) >= {"due_today", "due_soon", "classes", "reminders", "date"}
    no_controls = TestClient(create_app(Runtime(ctx, bus, ApprovalBroker(bus), lambda t: t, {})), client=LOCAL)
    assert no_controls.post("/api/control/listen").status_code == 409        # text mode: no mic to control


# ---------- meeting & lecture notes ----------

def speechy(seconds, amp=0.2, seed=0):
    rng = np.random.default_rng(seed)
    return (rng.normal(0, amp, int(seconds * 16000))).astype(np.float32)


def test_segmenter_cuts_near_pauses():
    from max_assistant.notes import Segmenter

    seg = Segmenter(min_s=25, max_s=32)
    audio = np.concatenate([speechy(24), np.zeros(8000, np.float32), speechy(12)])     # pause at 24 s
    chunks = seg.push(audio) + seg.flush()
    assert len(chunks) == 2
    assert abs(len(chunks[0][1]) / 16000 - 24.0) < 0.6          # cut inside the pause, not mid-"word"
    assert chunks[1][0] == len(chunks[0][1]) / 16000             # second chunk's timestamp follows on


def test_note_session_map_reduce_and_deadline_quotes(tmp_path):
    from max_assistant.notes import NoteSession, save

    said = iter(["Today we cover knowledge graphs and RDF triples.",
                 "Assignment two is due next Friday at 11:59 PM on Canvas.",
                 "Also there is a quiz on Monday about SPARQL."])
    calls = []

    def summarize(system, text, max_tokens):
        calls.append(system.split()[0])
        return "## Summary\nKnowledge graphs.\n## Key points\n- RDF" if "Combine" in system else f"- notes on: {text[:30]}"

    session = NoteSession("lecture", lambda audio, prompt: next(said), summarize, title="CSE 573", section_words=8)
    for _ in range(3):
        session.push(speechy(26))
        session.push(np.zeros(16000, np.float32))
    result = session.finish()
    assert result.word_count == sum(len(s.split()) for s in ["Today we cover knowledge graphs and RDF triples.",
                                                              "Assignment two is due next Friday at 11:59 PM on Canvas.",
                                                              "Also there is a quiz on Monday about SPARQL."])
    assert calls.count("You") >= 2 and calls[-1] == "Combine"   # section notes, then one final combine
    assert "## Deadlines & dates (exact quotes)" in result.notes_md
    assert "due next Friday at 11:59 PM" in result.notes_md and "quiz on Monday" in result.notes_md
    folder = save(result, tmp_path)
    assert (folder / "notes.md").read_text(encoding="utf-8").startswith("# CSE 573")
    assert "[00:" in (folder / "transcript.md").read_text(encoding="utf-8")


def test_note_session_skips_silence():
    from max_assistant.notes import NoteSession

    heard = []
    session = NoteSession("meeting", lambda a, p: heard.append(1) or "x", lambda s, t, n: "", section_words=100)
    session.push(np.zeros(16000 * 30, np.float32))
    result = session.finish()
    assert heard == [] and "Nothing was transcribed" in result.notes_md


def test_direct_routes_for_notes_and_forget():
    from max_assistant.agent import direct_route

    tools = {"forget", "start_notes", "stop_notes"}
    assert direct_route("take notes on this lecture", tools) == ("start_notes", {"kind": "lecture"})
    assert direct_route("I'm in a Teams call, can you take notes", tools) == ("start_notes", {"kind": "meeting"})
    assert direct_route("stop taking notes", tools) == ("stop_notes", {})
    assert direct_route("the meeting is over", tools) == ("stop_notes", {})
    assert direct_route("take a note: buy eggs", tools) is None          # quick note: not a recording
    assert direct_route("what were the key points of the lecture", tools) is None
    assert direct_route("take notes", {"forget"}) is None                # tool not available


def test_api_notes_start_stop():
    from fastapi.testclient import TestClient

    from max_assistant.events import ApprovalBroker, EventBus
    from max_assistant.server import Runtime, create_app

    class FakeNotes:
        def __init__(self):
            self.started = None

        def start(self, kind, title):
            self.started = (kind, title)
            return type("S", (), {"title": title or "Lecture", "kind": kind})()

        def stop(self):
            pass

        def status(self):
            return {"active": self.started is not None}

    ctx = digest_ctx()
    ctx.notes = FakeNotes()
    bus = EventBus()
    client = TestClient(create_app(Runtime(ctx, bus, ApprovalBroker(bus), lambda t: t, {})), client=LOCAL)
    r = client.post("/api/notes/start", json={"kind": "meeting", "title": "Standup"})
    assert r.status_code == 200 and r.json()["kind"] == "meeting" and ctx.notes.started == ("meeting", "Standup")
    assert client.get("/api/notes/status").json() == {"active": True}
    assert client.post("/api/notes/stop").json() == {"ok": True}


# ---------- phone access (Phase 4) ----------

def test_phone_requests_need_the_token_and_local_ones_dont():
    from max_assistant.remote import is_local

    assert is_local("127.0.0.1", {}) and is_local("::1", {})
    assert not is_local("127.0.0.1", {"x-forwarded-for": "100.64.0.7"})        # via tailscale serve
    assert not is_local("100.64.0.7", {})
    token = "t" * 40
    phone, rt = api_client(token=token, client=("100.64.0.7", 40000))
    assert phone.get("/api/state").status_code == 401
    assert phone.get("/api/state", headers={"Authorization": "Bearer wrong"}).status_code == 401
    auth = {"Authorization": f"Bearer {token}"}
    assert phone.get("/api/ping", headers=auth).json()["local"] is False
    assert phone.get("/api/pair", headers=auth).status_code == 403            # token never leaves the laptop
    assert phone.get("/api/state", params={"token": token}).status_code == 200
    with pytest.raises(Exception):
        with phone.websocket_connect("/api/ws") as ws:
            ws.receive_json()
    with phone.websocket_connect(f"/api/ws?token={token}") as ws:
        assert ws.receive_json()["kind"] == "hello"
    proxied, _ = api_client(token=token)
    assert proxied.get("/api/facts", headers={"X-Forwarded-For": "100.64.0.7"}).status_code == 401
    no_token, _ = api_client(token="", client=("100.64.0.7", 40000))
    assert no_token.get("/api/state", headers={"Authorization": "Bearer "}).status_code == 401


def test_phone_commands_are_tagged_and_pairing_gives_a_qr(monkeypatch):
    from max_assistant import remote

    seen = []
    token = "p" * 40
    run = lambda text, source="dashboard": seen.append(source) or "done"
    phone, rt = api_client(run, token=token, client=("100.64.0.7", 40000))
    auth = {"Authorization": f"Bearer {token}"}
    phone.post("/api/command", json={"text": "open youtube"}, headers=auth)
    laptop, _ = api_client(run, token=token)
    laptop.post("/api/command", json={"text": "open youtube"})
    assert seen == ["phone", "dashboard"]
    a = rt.approvals.open("Send it?")
    assert phone.post(f"/api/approvals/{a.id}", json={"approved": True}, headers=auth).json() == {"ok": True}
    assert a.by == "phone"
    monkeypatch.setattr(remote, "tailscale_url", lambda: "https://laptop.tail1234.ts.net")
    laptop, _ = api_client(token=token)
    p = laptop.get("/api/pair").json()
    assert p["url"] == "https://laptop.tail1234.ts.net" and p["token"] == token
    assert p["link"].startswith("max://pair?url=https%3A%2F%2Flaptop") and p["qr_svg"].lstrip().startswith("<")


def test_phone_voice_endpoint_transcribes_and_runs(monkeypatch):
    from max_assistant.remote import decode_wav, encode_wav

    t = np.arange(0, 1.0, 1 / 44100)
    wav = encode_wav((0.3 * np.sin(2 * np.pi * 440 * t)).astype(np.float32), 44100)
    audio = decode_wav(wav)
    assert abs(len(audio) - 16000) <= 1 and audio.dtype == np.float32
    seen = []
    client, rt = api_client(lambda text, source="dashboard": seen.append((text, source)) or "Opening YouTube.")

    class FakeSTT:
        def whisper(self, a, prompt=None):
            assert abs(len(a) - 16000) <= 1
            return "Hey Max, open YouTube."

    rt.ctx.stt = FakeSTT()
    r = client.post("/api/voice", content=wav, headers={"Content-Type": "audio/wav"}).json()
    assert r == {"heard": "open YouTube.", "reply": "Opening YouTube."} and seen == [("open YouTube.", "dashboard")]
    rt.ctx.stt.whisper = lambda a, prompt=None: "Thank you."
    assert client.post("/api/voice", content=wav).json() == {"heard": "", "reply": ""}
    assert client.post("/api/voice", content=b"not audio").status_code == 400


def test_phone_commands_need_a_tap_not_the_laptop_mic():
    import threading as th

    from max_assistant.events import ApprovalBroker, EventBus, current_origin, set_origin

    bus = EventBus()
    approvals = ApprovalBroker(bus)
    assert approvals.wait("Delete it?", timeout=0.05) is False               # no tap = no
    assert bus.recent[-1]["data"]["by"] == "timeout"
    th.Timer(0.1, lambda: approvals.answer(approvals.pending()[0]["id"], True, "phone")).start()
    assert approvals.wait("Delete it?", timeout=5) is True
    set_origin("phone")
    assert current_origin() == "phone"
    seen = []
    th.Thread(target=lambda: seen.append(current_origin())).start()
    time.sleep(0.05)
    assert seen == ["local"]                                                  # per thread
    set_origin("local")


def test_notes_summary_is_made_once_and_keeps_exact_deadline_quotes(tmp_path):
    from max_assistant.notes import NotesManager, make_summary, section

    notes_md = ("# Bio 101\n\n## Summary\nCells.\n\n## Key points\n- Mitochondria make ATP\n\n"
                "## Deadlines & dates (exact quotes)\n- [12:03] “Lab report due Friday at 11:59 PM.”\n")
    assert section(notes_md, "Key points") == "- Mitochondria make ATP"
    seen = []

    def summarize(system, text, max_tokens):
        seen.append(text)
        return "## TL;DR\nCells make energy.\n\n## Key takeaways\n- ATP\n\n## To do\n- Lab report due Friday at 5:59 PM"

    out = make_summary(summarize, notes_md)
    assert "exact quotes" not in seen[0]                       # the model never sees (or rewrites) the quotes
    assert out.endswith("## Exact quotes about dates\n- [12:03] “Lab report due Friday at 11:59 PM.”")

    folder = tmp_path / "2026-09-30 Bio 101"
    folder.mkdir()
    (folder / "notes.md").write_text(notes_md, encoding="utf-8")
    ctx = digest_ctx()
    ctx.memory = memory_store()
    mgr = NotesManager(ctx, lambda: None, tmp_path)
    mgr.summarize = summarize
    ctx.notes = mgr
    nid = ctx.memory.add_notes("Bio 101", "lecture", "2026-09-30T10:00", "2026-09-30T11:00", str(folder), "Cells.", 900)
    client, rt = api_client()
    rt.ctx.memory, rt.ctx.notes = ctx.memory, mgr
    assert client.get(f"/api/notes/{nid}").json()["summary_md"] == ""
    first = client.post(f"/api/notes/{nid}/summary").json()["summary_md"]
    assert "## TL;DR" in first and (folder / "summary.md").exists() and len(seen) == 2
    assert client.post(f"/api/notes/{nid}/summary").json()["summary_md"].strip() == first   # cached
    assert len(seen) == 2
    client.post(f"/api/notes/{nid}/summary", params={"refresh": True})
    assert len(seen) == 3
    assert client.get(f"/api/notes/{nid}").json()["summary_md"].startswith("## TL;DR")
    assert client.post("/api/notes/999/summary").status_code == 404


# ---------- live captions ----------

class FakeOnlineRecognizer:
    """Stands in for sherpa_onnx.OnlineRecognizer: 'hears' one word per 0.5 s of loud audio and
    ends a line after 1 s of quiet."""

    def __init__(self):
        self.words = iter("today we cover knowledge graphs and rdf triples assignment two is due friday".split())

    class Stream(dict):
        def accept_waveform(self, sr, x):
            self["last"] = x

    def create_stream(self):
        return self.Stream(text=[], loud=0, quiet=0)

    def is_ready(self, s):
        return False

    def decode_stream(self, s):
        pass

    def get_result(self, s):
        x = s["last"]
        blocks = round(len(x) / 1600)      # 100 ms blocks
        if float(np.sqrt(np.mean(x ** 2))) > 0.01:
            s["loud"] += blocks
            s["quiet"] = 0
            while s["loud"] >= 5:
                s["loud"] -= 5
                s["text"].append(next(self.words, "x"))
        else:
            s["quiet"] += blocks
        return " ".join(s["text"])

    def is_endpoint(self, s):
        return s["quiet"] >= 10 and bool(s["text"])

    def reset(self, s):
        s.update(text=[], loud=0, quiet=0)


def fake_live():
    from max_assistant.live import LiveModel, LiveTranscriber

    model = LiveModel.__new__(LiveModel)
    model.recognizer = FakeOnlineRecognizer()
    return LiveTranscriber(model)


def wait_for(cond, timeout=5.0):
    end = time.time() + timeout
    while time.time() < end and not cond():
        time.sleep(0.02)
    return cond()


def test_live_captions_show_words_then_lines():
    live = fake_live()
    tone = speechy(1.0)
    for i in range(0, len(tone), 1600):
        live.push(tone[i:i + 1600])                      # 100 ms blocks, like the mic
    assert wait_for(lambda: live.view(0)[1] == "today we")
    for _ in range(12):
        live.push(np.zeros(1600, np.float32))            # a pause ends the line
    assert wait_for(lambda: len(live.view(0)[0]) == 1)
    lines, partial = live.view(0)
    assert lines[0].text == "today we" and partial == "" and lines[0].start_s < 0.5
    assert live.view(after_s=lines[0].end_s)[0] == []    # once Whisper has covered it, it's dropped
    live.close()


def test_whisper_text_replaces_live_captions():
    from max_assistant.notes import NoteSession

    live = fake_live()
    release = threading.Event()

    def transcribe(audio, prompt):
        release.wait(5)                                  # Whisper is slow: hold it back
        return "Today we cover knowledge graphs."

    session = NoteSession("lecture", transcribe, lambda s, t, n: "- notes", title="CSE 573", live=live)
    block = speechy(26)
    for i in range(0, len(block), 1600):
        session.push(block[i:i + 1600])
    for _ in range(12):
        session.push(np.zeros(1600, np.float32))
    assert wait_for(lambda: session.live_view()["live"] != [])
    view = session.live_view()
    assert view["final"] == [] and view["captions"] and view["live"][0]["text"].startswith("today we cover")
    assert session.status()["last"].startswith("today we cover")        # overlay/phone see live words
    session.push(np.zeros(16000 * 8, np.float32))         # long enough for the segmenter to cut a chunk
    release.set()
    assert wait_for(lambda: session.live_view()["final"] != [])
    view = session.live_view()
    assert view["final"][0]["text"] == "Today we cover knowledge graphs." and view["live"] == []
    session.finish()


def test_live_captions_real_model():
    from max_assistant.config import ROOT
    from max_assistant.live import LiveModel, LiveTranscriber

    d = ROOT / "models/live/sherpa-onnx-streaming-zipformer-en-kroko-2025-08-06"
    wav = d / "test_wavs" / "0.wav"
    if not wav.exists():
        pytest.skip("live caption model not downloaded")
    import wave

    with wave.open(str(wav)) as w:
        x = np.frombuffer(w.readframes(w.getnframes()), np.int16).astype(np.float32) / 32768
    live = LiveTranscriber(LiveModel(d))
    for i in range(0, len(x), 1600):
        live.push(x[i:i + 1600])
    live.push(np.zeros(16000 * 3, np.float32))
    assert wait_for(lambda: any("country" in l.text.lower() for l in live.view(0)[0]) or "country" in live.view(0)[1].lower())
    live.close()


def test_config_has_no_duplicate_keys():
    """YAML silently keeps only the last of two same-named sections (a second 'notes:' once
    hid every meeting-notes setting)."""
    import yaml

    from max_assistant.config import ROOT

    class Strict(yaml.SafeLoader):
        pass

    def mapping(loader, node, deep=False):
        keys = [loader.construct_object(k, deep=deep) for k, _ in node.value]
        dupes = {k for k in keys if keys.count(k) > 1}
        assert not dupes, f"duplicate keys in config.yaml: {dupes}"
        return yaml.SafeLoader.construct_mapping(loader, node, deep)

    Strict.add_constructor(yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG, mapping)
    cfg = yaml.load((ROOT / "config.yaml").read_text(encoding="utf-8"), Loader=Strict)
    assert cfg["notes"]["file"] and cfg["notes"]["live_model"]


# ---------- commands from the phone open things on the phone ----------

def test_phone_searches_open_on_the_phone_not_the_laptop(tmp_path, monkeypatch):
    from max_assistant.events import EventBus, set_origin

    opened = []
    monkeypatch.setattr(system_tools.webbrowser, "open", opened.append)
    reg = make_registry(tmp_path)
    reg.context.bus = EventBus()
    set_origin("phone")
    try:
        out = reg.run("open_website", {"target": "cute dog photos", "section": "images"})
    finally:
        set_origin("local")
    assert out == "Here are Google Images results for cute dog photos on your phone."
    assert opened == []                                                       # laptop untouched
    ev = reg.context.bus.recent[-1]
    assert ev["kind"] == "phone_open" and ev["data"]["url"] == "https://www.google.com/search?q=cute+dog+photos&udm=2"
    assert reg.run("open_website", {"target": "youtube.com"}) == "Opened youtube.com."   # laptop commands unchanged
    assert opened == ["https://youtube.com"]


def test_phone_command_redirects_laptop_browser_tools_and_fixes_empty_promises(tmp_path, monkeypatch):
    """The real failure: from the phone Max called browser_tabs (no page open), then said
    "I'll open a search for cute dog photos" without opening anything."""
    from max_assistant.events import EventBus, set_origin

    monkeypatch.setattr(system_tools.webbrowser, "open", lambda url: None)
    reg = make_registry(tmp_path)
    reg.context.bus = EventBus()

    @reg.tool("List or switch the browser's tabs.", params={})
    def browser_tabs():
        raise AssertionError("must not touch the laptop's browser for a phone command")

    llm = FakeLLM([ChatReply("", [ToolCall("browser_tabs", {})]),
                   ChatReply("I'll open a search for cute dog photos right away."),        # empty promise
                   ChatReply("", [ToolCall("open_website", {"target": "cute dog photos", "section": "images"})]),
                   ChatReply("Here are cute dog photos on your phone.")])     # "search ... and show me" = two steps
    agent = Agent(llm, reg, "fast", "planner")
    set_origin("phone")
    try:
        answer = agent.handle("can you search for some cute dog photos and show me?")
    finally:
        set_origin("local")
    assert "on your phone" in answer
    assert "can't see the laptop's browser" in llm.calls[1]["messages"][-1]["content"]
    assert "System check" in llm.calls[2]["messages"][-1]["content"]
    assert reg.context.bus.recent[-1]["kind"] == "phone_open"


# ---------- Max acting on the phone (phone.py, tools/phone.py) ----------

def phone_setup(tmp_path, answer=None):
    """A registry with phone tools and a fake phone that answers every action."""
    from max_assistant.events import EventBus
    from max_assistant.phone import PhoneBridge
    from max_assistant.tools import phone as phone_tools

    bus = EventBus()
    reg = make_registry(tmp_path)
    reg.context.phone = bridge = PhoneBridge(bus, timeout=3)
    phone_tools.register(reg)
    sent = []
    if answer is not None:
        bridge.connected(+1)
        real_publish = bus.publish

        def publish(kind, data=None):
            ev = real_publish(kind, data)
            if kind == "phone_action":
                sent.append(data)
                threading.Thread(target=lambda: bridge.resolve(data["id"], answer(data))).start()
            return ev
        bus.publish = publish
    return reg, bridge, sent


def test_phone_tools_send_actions_and_speak_the_phones_answer(tmp_path):
    reg, bridge, sent = phone_setup(tmp_path, lambda d: {"ok": True, "message": f"did {d['action']}"})
    bridge.set_state({"apps": [{"label": "Instagram", "package": "com.instagram.android"},
                               {"label": "WhatsApp", "package": "com.whatsapp"}]})
    assert reg.run("phone_open_app", {"name": "insta"}) == "did open_app"
    assert sent[-1]["params"] == {"name": "insta", "package": "com.instagram.android"}
    assert reg.run("phone_timer", {"duration": "an hour and 30 minutes", "label": "pasta"}) == "did timer"
    assert sent[-1]["params"] == {"seconds": 5400, "label": "pasta"}
    reg.run("phone_alarm", {"time": "tomorrow at 6:45am"})
    assert (sent[-1]["params"]["hour"], sent[-1]["params"]["minute"]) == (6, 45)
    reg.run("phone_message", {"contact": "Mom", "text": "running late", "app": "telegram"})
    assert sent[-1]["params"] == {"contact": "Mom", "text": "running late", "app": "sms"}
    assert "phone_calendar_add" not in reg.tools          # events go to Max's calendar (add_event) now
    assert reg.tools["phone_call"].risky and not reg.tools["phone_message"].risky   # texts: the user taps Send
    reg.run("phone_notifications", {"app": "all"})
    assert sent[-1]["params"] == {"app": "", "count": 8}


def test_phone_tools_when_the_phone_is_away(tmp_path):
    from max_assistant.tools.phone import parse_duration

    reg, bridge, _ = phone_setup(tmp_path)
    assert reg.run("phone_status", {}) == "Your phone isn't connected to Max right now."
    bridge.connected(+1)
    bridge.timeout = 0.05
    assert "didn't answer in time" in reg.run("phone_status", {})
    assert reg.run("phone_timer", {"duration": "soon"}).startswith("Error")
    assert parse_duration("10 minutes") == 600 and parse_duration("90s") == 90 and parse_duration("half an hour") == 1800
    assert parse_duration("5") == 300


def test_phone_api_results_state_and_connection_count():
    from max_assistant.events import EventBus
    from max_assistant.phone import PhoneBridge

    token = "q" * 40
    phone, rt = api_client(token=token, client=("100.64.0.7", 40000))
    rt.ctx.phone = bridge = PhoneBridge(rt.bus, timeout=3)
    auth = {"Authorization": f"Bearer {token}"}
    with phone.websocket_connect(f"/api/ws?token={token}") as ws:
        ws.receive_json()
        assert bridge.online
        assert phone.post("/api/phone/state", json={"apps": [{"label": "Maps", "package": "com.google.maps"}], "battery": 80},
                          headers=auth).json() == {"ok": True, "apps": 1}
        assert bridge.apps() == {"maps": "com.google.maps"}
        out = {}
        t = threading.Thread(target=lambda: out.update(bridge.request("status")))
        t.start()
        while True:                                               # the phone hears the action over the socket
            ev = ws.receive_json()
            if ev["kind"] == "phone_action":
                break
        assert phone.post("/api/phone/result", json={"id": ev["data"]["id"], "ok": True, "message": "Battery 80%."},
                          headers=auth).json() == {"ok": True}
        t.join(3)
        assert out == {"ok": True, "message": "Battery 80%."}
    assert not bridge.online
    assert phone.post("/api/phone/result", json={"id": 999, "message": "late"}, headers=auth).status_code == 404


# ---------- speaking while thinking, and interrupting (speech.py) ----------

def test_sentences_are_spoken_as_they_arrive_and_the_rest_at_the_end():
    from max_assistant.speech import SentenceSpeaker

    said = []
    s = SentenceSpeaker(lambda t: said.append(t))
    s.feed(None)
    for piece in ["Two things are due", " this week. The project", " report is Friday at 11:59 PM.", " Also a qu", "iz on Monday"]:
        s.feed(piece)
    s._q.join()
    assert said == ["Two things are due this week.", "The project report is Friday at 11:59 PM."]   # spoken early
    assert s.finish("Two things are due this week. The project report is Friday at 11:59 PM. Also a quiz on Monday") is None
    assert said[-1] == "Also a quiz on Monday"                                            # only the rest


def test_a_claimed_action_halts_streaming_so_the_agent_can_correct_it():
    from max_assistant.speech import SentenceSpeaker

    said = []
    s = SentenceSpeaker(lambda t: said.append(t))
    s.feed(None)
    s.feed("Sure. I'll remind you tomorrow. Anything else? ")
    s._q.join()
    assert said == ["Sure."]
    s.feed(None)                                     # the nudged second attempt
    s.feed("Okay, I'll remind you tomorrow at 7 AM.")
    s.finish("Okay, I'll remind you tomorrow at 7 AM.")
    assert said[-1] == "Okay, I'll remind you tomorrow at 7 AM."


def test_interrupting_stops_the_remaining_sentences():
    from max_assistant.speech import SentenceSpeaker

    said = []
    s = SentenceSpeaker(lambda t: said.append(t) or ("stop" if len(said) == 1 else None))
    s.feed(None)
    s.feed("First sentence. Second sentence. Third ")
    assert s.finish("First sentence. Second sentence. Third sentence.") == "stop"
    assert said == ["First sentence."]


def test_watch_for_interrupt_ignores_maxs_own_words():
    from max_assistant.speech import watch_for_interrupt

    frames = iter(["a", "b", "c", "d"])
    heard = iter([None, "STOP", None, "HEY_MAX"])
    read = lambda timeout: next(frames, None)
    assert watch_for_interrupt(read, lambda f: next(heard), 5, "Take the bus to the next stop.") == "wake"
    frames, heard = iter(["a"]), iter(["MAX_STOP"])
    assert watch_for_interrupt(lambda t: next(frames, None), lambda f: next(heard), 5, "Here you go.") == "stop"
    assert watch_for_interrupt(lambda t: None, lambda f: None, 0.05, "Quiet.") is None


def test_converse_after_an_interrupt():
    from max_assistant.main import converse

    replies = iter(["Here's a long answer?", "Sure."])
    cut = iter(["wake", None])
    log = []
    n = converse("tell me about rdf", lambda t: log.append(t) or next(replies), lambda: "should not be used",
                 interrupted=lambda: next(cut), listen_new=lambda: "actually what's the weather")
    assert n == 2 and log == ["tell me about rdf", "actually what's the weather"]
    n = converse("tell me about rdf", lambda t: "Is that all?", lambda: "unused", interrupted=lambda: "stop")
    assert n == 1


def test_llm_streams_text_and_still_returns_tool_calls(monkeypatch):
    import json as _json

    from max_assistant import llm as L

    lines = [{"message": {"content": "Hello"}}, {"message": {"content": " there."}},
             {"message": {"content": "", "tool_calls": [{"function": {"name": "web_search", "arguments": {"query": "x"}}}]}},
             {"done": True, "message": {"content": ""}}]

    class R:
        status_code = 200
        text = ""
        def iter_lines(self):
            return [_json.dumps(l).encode() for l in lines]

    monkeypatch.setattr(L.requests, "post", lambda url, json, timeout, **kw: R())
    pieces = []
    reply = L.OllamaClient("http://x").chat("m", [], on_text=pieces.append)
    assert pieces == ["Hello", " there."] and reply.content == "Hello there."
    assert reply.tool_calls[0].name == "web_search"


def test_agent_does_not_stream_while_a_requested_tool_is_pending():
    reg, ctx = memory_registry()
    agent = Agent(FakeLLM([]), reg, "fast", "planner")
    assert agent._stream_for(print, "remind me about gym tomorrow", []) is None
    assert agent._stream_for(print, "what is rdf", []) is print
    assert agent._stream_for(None, "what is rdf", []) is None


# ---------- deadline countdown (alerts.py) ----------

def test_deadline_alerts_fire_once_per_window_and_for_classes():
    import datetime as dt

    from max_assistant.alerts import DeadlineAlerts
    from max_assistant.canvas import CanvasItem

    now = dt.datetime(2026, 10, 1, 9, 0)
    lab = CanvasItem("assignment", "Lab 4", "CSE 572", dt.datetime(2026, 10, 1, 23, 59), True, url="https://canvas/lab4")
    quiz = CanvasItem("assignment", "Quiz 3", "CSE 579", dt.datetime(2026, 10, 2, 10, 0), False)
    cls = CanvasItem("class", "Semantic Web Mining", "CSE 573", dt.datetime(2026, 10, 1, 9, 8), False, link="https://asu.zoom.us/j/1")

    class Feed:
        def between(self, start, end, kind=None):
            return [i for i in (lab, quiz, cls) if start <= i.start < end]

    ctx = Context(load_config(), FakeLLM([]))
    ctx.memory, ctx.canvas = memory_store(), Feed()
    events = []
    alerts = DeadlineAlerts(ctx, lambda kind, data: events.append(data))
    out = alerts.check(now)
    titles = sorted(a["title"] for a in out)
    assert titles == ["Class in 8 minutes: Semantic Web Mining", "Due in 15 hours: Lab 4"]
    cls_alert = next(a for a in out if a["kind"] == "class")
    assert cls_alert["url"] == "https://asu.zoom.us/j/1" and "tap to join" in cls_alert["text"]
    assert alerts.check(now + dt.timedelta(minutes=1)) == []                       # once only
    later = alerts.check(dt.datetime(2026, 10, 1, 21, 30))                          # 2.5 h before the lab
    assert sorted(a["title"] for a in later) == ["Due in 12 hours: Quiz 3", "Due in 2 hours: Lab 4"]
    late = DeadlineAlerts(ctx, lambda k, d: None)                                   # laptop was off: only the 3 h alert
    ctx.memory.set("alerts_sent", "[]")
    assert [a["title"] for a in late.check(dt.datetime(2026, 10, 2, 8, 0)) if "Quiz" in a["title"]] == ["Due in 2 hours: Quiz 3"]


# ---------- privacy page ----------

def test_privacy_page_lists_exports_and_deletes(tmp_path):
    import io
    import zipfile

    client, rt = api_client()
    cfg = rt.ctx.cfg
    cfg["notes"] = {**(cfg.get("notes") or {}), "folder": str(tmp_path / "Max Notes"), "file": str(tmp_path / "quick.md")}
    cfg["browser"] = {**(cfg.get("browser") or {}), "profile_dir": str(tmp_path / "profile")}
    m = rt.ctx.memory
    m.add_fact("Manas's exam is on Friday")
    m.log_turn("hi", "Hello!")
    folder = tmp_path / "Max Notes" / "2026-10-01 Lecture"
    folder.mkdir(parents=True)
    (folder / "notes.md").write_text("# Lecture\n- RDF", encoding="utf-8")
    m.add_notes("Lecture", "lecture", "2026-10-01T10:00", "2026-10-01T11:00", str(folder), "RDF", 10)
    (tmp_path / "quick.md").write_text("- buy milk", encoding="utf-8")
    (tmp_path / "profile").mkdir()
    (tmp_path / "profile" / "Cookies").write_bytes(b"x" * 10)

    cats = {c["id"]: c for c in client.get("/api/privacy").json()}
    assert cats["facts"]["amount"] == "1 facts" and cats["conversations"]["amount"] == "1 turns"
    z = zipfile.ZipFile(io.BytesIO(client.get("/api/privacy/export").content))
    names = z.namelist()
    assert "facts.json" in names and "notes/2026-10-01 Lecture/notes.md" in names and "quick_notes.md" in names
    assert not any("Cookies" in n or "token" in n for n in names)                   # secrets stay out
    assert "Friday" in z.read("facts.json").decode()

    assert client.post("/api/privacy/delete/facts", json={"confirm": "nope"}).status_code == 400
    assert client.post("/api/privacy/delete/facts", json={"confirm": "facts"}).json()["ok"]
    assert m.count("facts") == 0 and m.search_facts("exam") == []
    client.post("/api/privacy/delete/notes", json={"confirm": "notes"})
    assert not folder.exists() and m.count("notes") == 0
    client.post("/api/privacy/delete/browser", json={"confirm": "browser"})
    assert not (tmp_path / "profile").exists()
    assert client.post("/api/privacy/delete/bogus", json={"confirm": "bogus"}).status_code == 404
    phone, _ = api_client(token="t" * 40, client=("100.64.0.7", 40000))
    assert phone.get("/api/privacy", headers={"Authorization": "Bearer " + "t" * 40}).status_code == 403   # never over the phone link


# ---------- ask your course material (course.py) ----------

def make_course_files(root):
    import docx
    import pymupdf
    from pptx import Presentation

    (root / "CSE 573").mkdir(parents=True)
    (root / "CSE 572").mkdir()
    pdf = pymupdf.open()
    for text in ["Week 1 overview of the semantic web course.",
                 "SPARQL OPTIONAL keeps results even when the optional pattern has no match, leaving the variable unbound."]:
        pdf.new_page().insert_text((72, 72), text)
    pdf.save(str(root / "CSE 573" / "Lab 3 handout.pdf"))
    pres = Presentation()
    s = pres.slides.add_slide(pres.slide_layouts[1])
    s.shapes.title.text = "RDF Schema"
    s.placeholders[1].text = "RDFS adds classes, subclasses and domain and range constraints to RDF vocabularies."
    pres.save(str(root / "CSE 573" / "Lecture 5.pptx"))
    d = docx.Document()
    d.add_paragraph("The midterm covers decision trees, entropy and information gain, and k-means clustering.")
    d.save(str(root / "CSE 572" / "Midterm review.docx"))
    (root / "CSE 572" / "~$Midterm review.docx").write_text("lock file", encoding="utf-8")   # Word's temp file: skipped


def test_course_library_indexes_files_and_answers_with_sources(tmp_path):
    from max_assistant.course import CourseLibrary

    make_course_files(tmp_path / "Course")
    lib = CourseLibrary([tmp_path / "Course"], WordEmbedder(), tmp_path / "course.db")
    assert lib.scan()["indexed"] == 3
    hit = lib.search("what does SPARQL OPTIONAL do when there is no match")[0]
    assert (hit.file, hit.page, hit.course) == ("Lab 3 handout.pdf", "p. 2", "CSE 573")
    assert lib.search("rdfs domain range constraints")[0].page == "slide 1"
    assert lib.search("midterm entropy information gain", course="CSE 572")[0].file == "Midterm review.docx"
    assert lib.scan()["indexed"] == 0                                  # unchanged files aren't re-read
    (tmp_path / "Course" / "CSE 572" / "Midterm review.docx").unlink()
    assert lib.scan()["removed"] == 1
    assert lib.stats()["courses"] == {"CSE 573": 2}
    lib.clear()
    assert lib.search("SPARQL") == [] and lib.stats()["passages"] == 0


def test_course_tools_give_the_model_cited_passages(tmp_path):
    from max_assistant.course import CourseLibrary
    from max_assistant.tools import course as course_tools

    make_course_files(tmp_path / "Course")
    reg = make_registry(tmp_path)
    reg.context.course = lib = CourseLibrary([tmp_path / "Course"], WordEmbedder(), tmp_path / "course.db")
    course_tools.register(reg)
    assert reg.run("course_files", {}).startswith("No course files yet")
    lib.scan()
    out = reg.run("course_search", {"question": "SPARQL OPTIONAL no match"})
    assert "[Lab 3 handout.pdf, p. 2]" in out and "say where it's from" in out
    assert reg.run("course_files", {}) == "I have 3 files: 1 for CSE 572, 2 for CSE 573."
    assert reg.run("course_search", {"question": "photosynthesis in plants"}).startswith("Nothing in the course files")


# ---------- only the tools a request needs (toolselect.py) ----------

def test_tool_selection_keeps_requests_small_and_relevant(tmp_path):
    from max_assistant.events import EventBus
    from max_assistant.phone import PhoneBridge
    from max_assistant.tools import phone as phone_tools
    from max_assistant.tools.memory import register as memory_register
    from max_assistant.toolselect import CORE, ToolSelector

    reg, ctx = memory_registry()
    system_tools.register(reg)
    ctx.phone = PhoneBridge(EventBus())
    phone_tools.register(reg)
    sel = ToolSelector(reg, WordEmbedder())
    total = len(reg.tools)
    names = sel.select("set a timer for 10 minutes on my phone")
    assert "phone_timer" in names and "phone_call" in names and len(names) < total   # the whole phone family
    assert names[:len([c for c in CORE if c in reg.tools])] == [c for c in CORE if c in reg.tools]
    assert "set_reminder" in sel.select("remind me to call mom at 6")
    assert "browser_click" not in sel.select("what's the weather")
    follow = sel.select("the second one", recent_tools=["phone_message"])
    assert "phone_message" in follow                                                   # follow-ups keep their tools
    assert len(ToolSelector(reg, WordEmbedder(), max_tools=8).select("remind me on my phone to remember things")) <= 15


def test_context_overflow_retries_without_old_conversation():
    from max_assistant.llm import OllamaError

    class Overflow(FakeLLM):
        def chat(self, model, messages, tools=None, **kw):
            if len(messages) > 3:
                self.calls.append({"messages": list(messages)})
                raise OllamaError("Ollama error 400: request (4315 tokens) exceeds the available context size")
            return super().chat(model, messages, tools)

    reg, ctx = memory_registry()
    llm = Overflow([ChatReply("Hi again.")])
    agent = Agent(llm, reg, "fast", "planner")
    agent.history = [[{"role": "user", "content": "earlier"}, {"role": "assistant", "content": "ok"}]]
    agent._last_turn = time.monotonic()
    assert agent.handle("hello") == "Hi again."


# ---------- read my screen (tools/screen.py) ----------

def test_read_screen_reads_text_and_keeps_nothing(tmp_path, monkeypatch):
    if sys.platform != "win32":
        pytest.skip("Windows OCR")
    from PIL import Image, ImageDraw, ImageFont

    from max_assistant.tools import screen as screen_tools

    img = Image.new("RGB", (1000, 220), "white")
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype("consola.ttf", 26)
    d.text((20, 40), "TypeError: unsupported operand type(s) for +: 'int' and 'str'", fill="black", font=font)
    d.text((20, 110), 'File "app.py", line 12, in total', fill="black", font=font)
    monkeypatch.setattr(screen_tools, "capture", lambda whole=False: (img, "app.py - VS Code"))   # never the real screen
    reg = make_registry(tmp_path)
    screen_tools.register(reg)
    out = reg.run("read_screen", {"question": "what does this error mean?"})
    assert "unsupported operand" in out and "app.py - VS Code" in out and "what does this error mean?" in out
    monkeypatch.setattr(screen_tools, "capture", lambda whole=False: (Image.new("RGB", (400, 200), "white"), "blank"))
    assert reg.run("read_screen", {"question": "?"}) == "I couldn't find any text in that window."


# ---------- command accuracy test (evals.py) ----------

def test_eval_cases_load_and_judging():
    from max_assistant.evals import judge, load_cases, summarize

    cases = load_cases()
    assert len(cases) >= 80 and all(t["expect"] for c in cases for t in c["turns"])
    turn = {"expect": ["phone_timer"], "args": {"duration": "10"}}
    assert judge(turn, [("phone_timer", {"duration": "10 minutes"})]) == (True, True, "")
    assert judge(turn, [("phone_timer", {"duration": "5 minutes"})])[:2] == (True, False)
    assert judge(turn, [("set_reminder", {})])[0] is False
    assert judge({"expect": ["none"], "args": {}}, []) == (True, True, "")
    s = summarize([{"category": "a", "tool_ok": True, "args_ok": False, "seconds": 1.0, "prompt_tokens": 100},
                   {"category": "a", "tool_ok": True, "args_ok": True, "seconds": 3.0, "prompt_tokens": 300}])
    assert s["tool_accuracy"] == 1.0 and s["full_accuracy"] == 0.5 and s["categories"]["a"] == {"n": 2, "tool": 2, "full": 1}


# ---------- voice ID (voiceid.py) ----------

def test_voice_id_tells_the_user_from_other_voices(tmp_path):
    from max_assistant.config import ROOT
    from max_assistant.tts import PiperTTS
    from max_assistant.voiceid import VoiceID

    model = ROOT / "models/voiceid/eres2net_en_voxceleb.onnx"
    voice = ROOT / "models/piper/en_US-ryan-medium.onnx"
    if not (model.exists() and voice.exists()):
        pytest.skip("voice ID model or Piper voice not downloaded")
    ryan = PiperTTS(str(voice))

    def say(line, pitch=1.0):
        a, sr = ryan.synthesize(line)
        t = np.arange(0, len(a) / sr, 1 / 16000)
        x = np.interp(t * pitch, np.arange(len(a)) / sr, a).astype(np.float32)   # resampled = a different-sounding voice
        return x

    vid = VoiceID(model, tmp_path / "profile.npy")
    assert vid.score(say("hello there")) is None and not vid.enrolled            # nothing enrolled yet
    vid.enroll([say(l) for l in ["What's due this week?", "Remind me to call mom tomorrow.", "Open Spotify please.",
                                 "Take notes for this lecture."]])
    assert vid.enrolled and (tmp_path / "profile.npy").exists()
    assert vid.is_owner(say("Shut down the laptop now.")) is True
    assert vid.score(np.zeros(4000, np.float32)) is None                          # too short to judge
    assert VoiceID(model, tmp_path / "profile.npy").enrolled                      # saved and reloaded
    vid.forget()
    assert not vid.enrolled and not (tmp_path / "profile.npy").exists()


def test_read_screen_drops_the_text_cursor_and_window_furniture():
    from max_assistant.tools.screen import clean, merge_reads

    # The real report: Notepad said "hey how are you" and Max read "hey how are youl"
    first = ["File Edit View", "hey how are youl", "Ln 1, Col 16 100% Windows (CRLF) UTF-8"]
    second = ["File Edit View", "hey how are you", "Ln 1, Col 16 100% Windows (CRLF) UTF-8"]
    assert clean(merge_reads(first, second)) == ["hey how are you"]
    assert clean(merge_reads(second, first)) == ["hey how are you"]        # cursor in the other read
    assert merge_reads(["I love it all"], ["I love it al"]) == ["I love it al"]   # (only one stray char is trimmed)
    assert clean(["Run the file", "Edit View Help"]) == ["Run the file"]   # real sentences with menu words stay


# ---------- natural voice with Piper fallback (tts.py) ----------

def test_natural_voice_trims_padding_and_falls_back_to_piper():
    import concurrent.futures as cf

    from max_assistant.tts import NaturalVoice, trim_silence

    sr = 24000
    speech = np.concatenate([np.zeros(sr), 0.3 * np.ones(sr // 2, np.float32), np.zeros(sr)]).astype(np.float32)
    trimmed = trim_silence(speech, sr)
    assert abs(len(trimmed) - (sr // 2 + 2 * int(0.06 * sr))) <= 2           # Ava's padding removed
    assert len(trim_silence(np.zeros(sr, np.float32), sr)) == 0

    class Piper:
        def synthesize(self, text):
            return np.ones(10, np.float32), 22050

    v = NaturalVoice.__new__(NaturalVoice)                                  # no real voice or network in tests
    v.name, v.fallback, v.timeout_s = "Microsoft Ava Online", Piper(), 1.0
    v._pool = cf.ThreadPoolExecutor(max_workers=1)
    v._speak = lambda text: (np.full(100, 0.5, np.float32), 24000)
    assert v.synthesize("hi")[1] == 24000                                   # online voice when it works
    v._speak = lambda text: (_ for _ in ()).throw(OSError("no internet"))
    assert v.synthesize("hi")[1] == 22050                                   # Piper when offline
    v._speak = lambda text: (np.zeros(0, np.float32), 24000)
    assert v.synthesize("hi")[1] == 22050                                   # Piper on silence


# ---------- Gmail (mail.py, tools/mail.py) ----------

RAW_MAIL = (b"From: Prof. Jane Smith <jsmith@asu.edu>\r\nTo: manas@example.com\r\nSubject: Midterm moved\r\n"
            b"Date: Thu, 01 Oct 2026 09:30:00 -0700\r\nMessage-ID: <abc123@asu.edu>\r\n"
            b"Content-Type: multipart/alternative; boundary=b1\r\n\r\n"
            b"--b1\r\nContent-Type: text/plain; charset=utf-8\r\n\r\nHi all,\nThe midterm is moved to Monday at 10am.\n\n"
            b"On Wed, Sep 30, 2026 at 9:00 AM Someone <x@y.com> wrote:\n> old quoted text\n"
            b"--b1\r\nContent-Type: text/html; charset=utf-8\r\n\r\n<p>Hi all,</p><p>The midterm is moved.</p>\r\n--b1--\r\n")


class FakeGmail:
    def __init__(self, mails):
        self.mails, self.sent, self.archived, self.read, self.queries = mails, [], [], [], []

    def search(self, query="is:unread in:inbox", limit=5, account=None):
        self.queries.append(query)
        return self.mails[:limit]

    def get(self, msgid, account=None):
        return next((m for m in self.mails if m.msgid == msgid), None)

    def send(self, to, subject, body, reply_to=None, account=None):
        self.sent.append((to, subject, body, reply_to.message_id if reply_to else None))

    def archive(self, msgid, account=None):
        self.archived.append(msgid)

    def mark_read(self, msgid, account=None):
        self.read.append(msgid)

    def contacts(self, name, limit=40):
        return [(m.sender, m.address) for m in self.mails if name.lower() in m.sender.lower()]


def test_mail_parsing_prefers_plain_text_and_drops_quotes():
    from max_assistant.mail import html_to_text, parse

    m = parse(RAW_MAIL, "111", "222", flags="", labels="\\Important \\Inbox")
    assert (m.sender, m.address, m.subject) == ("Prof. Jane Smith", "jsmith@asu.edu", "Midterm moved")
    assert "moved to Monday at 10am" in m.body and "old quoted text" not in m.body
    assert m.unread and m.important and m.message_id == "<abc123@asu.edu>" and m.link.endswith("/de")
    text = html_to_text("<p>Hi&amp;there</p><script>x()</script><br>Bye")
    assert "Hi&there" in text and "Bye" in text and "x()" not in text


def test_mail_tools_check_read_reply_archive(tmp_path):
    from max_assistant.mail import parse
    from max_assistant.tools import mail as mail_tools

    m1 = parse(RAW_MAIL, "111", "222", "", "\Important")
    m2 = parse(RAW_MAIL.replace(b"Prof. Jane Smith <jsmith@asu.edu>", b"Canvas <notifications@instructure.com>")
               .replace(b"Midterm moved", b"New grade posted"), "333", "444")
    reg = make_registry(tmp_path)
    reg.context.mail = gmail = FakeGmail([m1, m2])
    mail_tools.register(reg)
    out = reg.run("mail_check", {})
    assert "1. From Prof. Jane Smith" in out and "[important]" in out and "2. From Canvas" in out
    assert "moved to Monday" in reg.run("mail_read", {"which": "the first one"})
    assert "New grade" in reg.run("mail_read", {"which": "canvas"})
    reg.run("mail_reply", {"which": "1", "text": "Thanks, noted!"})
    assert gmail.sent == [("jsmith@asu.edu", "Re: Midterm moved", "Thanks, noted!", "<abc123@asu.edu>")]
    assert reg.tools["mail_reply"].risky and reg.tools["mail_send"].risky and not reg.tools["mail_archive"].risky
    assert reg.run("mail_archive", {"which": "second"}) == "Archived the email from Canvas." and gmail.archived == ["333"]
    assert reg.run("mail_mark_read", {"which": "all"}) == "Marked 2 emails as read."
    assert reg.run("mail_send", {"to": "jane", "body": "Hi"}) == "Sent your email to jane." and gmail.sent[-1][0] == "jsmith@asu.edu"
    gmail.mails = []
    assert reg.run("mail_check", {}) == "No unread email in your inbox."


def test_mail_alerts_for_chosen_senders_and_digest_section():
    from max_assistant.alerts import MailAlerts
    from max_assistant.digest import gather, spoken
    from max_assistant.mail import parse

    m1 = parse(RAW_MAIL, "111", "222", "", "\Important")
    ctx = Context(load_config(), FakeLLM([]))
    ctx.memory = memory_store()
    ctx.mail = FakeGmail([m1])
    events = []
    alerts = MailAlerts(ctx, lambda k, d: events.append(d), ["@asu.edu"])
    assert alerts.check() == []                                  # first run: existing mail doesn't alert
    m2 = parse(RAW_MAIL.replace(b"Midterm moved", b"Office hours today"), "555", "666")
    ctx.mail.mails = [m2, m1]
    assert [a["title"] for a in alerts.check()] == ["Email from Prof. Jane Smith"] and events[-1]["text"] == "Office hours today"
    assert alerts.check() == []                                  # once only
    restarted = MailAlerts(ctx, lambda k, d: None, ["@asu.edu"])
    ctx.mail.mails = [parse(RAW_MAIL.replace(b"Midterm moved", b"Grades out"), "777", "888"), m2, m1]
    assert len(restarted.check()) == 1                           # mail that came while Max was off still alerts
    ctx.mail.mails = [m1]
    plan = gather(ctx)
    assert plan.mail and "Important email: Prof. Jane Smith about Midterm moved." in spoken(plan)


def test_three_mailboxes_merge_and_reply_from_the_right_one():
    from max_assistant.mail import MailAccounts, accounts_from_secrets, parse

    assert accounts_from_secrets({"personal": {"address": "you@gmail.com", "app_password": "xxxx xxxx xxxx xxxx"}}) is None
    accts = accounts_from_secrets({"personal": {"address": "manas@example.com", "app_password": "abcd efgh ijkl mnop"},
                                   "university": {"address": "you@asu.edu", "app_password": "xxxx xxxx xxxx xxxx"},
                                   "max": {"address": "max.helper@example.com", "app_password": "qrst uvwx yzab cdef"}})
    assert accts.names == ["personal", "max"]                     # the untouched placeholder slot is skipped
    assert accounts_from_secrets({"address": "a@b.com", "app_password": "real pass word here"}).names == ["personal"]

    uni = parse(RAW_MAIL, "1", "2")
    home = parse(RAW_MAIL.replace(b"Thu, 01 Oct 2026 09:30", b"Thu, 01 Oct 2026 11:30")
                 .replace(b"Midterm moved", b"Dinner tonight?"), "3", "4")
    clients = {"personal": FakeGmail([home]), "university": FakeGmail([uni])}
    multi = MailAccounts(clients)
    found = multi.search()
    assert [(m.subject, m.account) for m in found] == [("Dinner tonight?", "personal"), ("Midterm moved", "university")]
    assert [m.account for m in multi.search(account="school")] == ["university"]
    multi.send("jsmith@asu.edu", "Re: Midterm moved", "Thanks", reply_to=found[1])
    assert clients["university"].sent and not clients["personal"].sent      # replies go out from where the mail came
    multi.send("x@y.com", "Hi", "Hello", account="uni")
    assert len(clients["university"].sent) == 2


# ---------- notes from course documents (docnotes.py) ----------

class FakeSummarizer:
    """Stands in for the model: records each call and returns recognisable notes."""

    def __init__(self):
        self.calls = []

    def __call__(self, system, text, max_tokens):
        self.calls.append((system, text))
        if "Combine" in system:
            return "## Summary\nNotes about " + " ".join(text.split()[:4]) + "\n\n## Key points\n- point"
        return f"- section notes {len(self.calls)}"


def test_document_sections_group_pages_and_split_long_ones():
    from max_assistant.docnotes import sections

    items = [("p. 1", "a " * 300), ("p. 2", "b " * 300), ("p. 3", "c " * 900)]
    out = sections(items, words=700)
    assert [w for w, _ in out] == ["p. 1-3", "p. 3", "p. 3"]
    assert [len(t.split()) for _, t in out] == [700, 700, 100]
    assert sections([("slide 1", "one two three")]) == [("slide 1", "one two three")]


def test_document_notes_read_the_whole_file_and_quote_deadlines(tmp_path):
    from max_assistant.docnotes import TEXT_FILE, save, write_notes

    f = tmp_path / "Syllabus.txt"
    f.write_text("# Intro\n" + "Semantic web basics and RDF triples. " * 150 +
                 "\n# Grading\nThe final project is due Friday, December 5 at 11:59 PM.\n", encoding="utf-8")
    fake = FakeSummarizer()
    r = write_notes(f, fake, course="CSE 573")
    assert len(fake.calls) >= 3                                    # map per section, then combine
    assert r.notes_md.startswith("## Summary")
    assert "[section 2] “The final project is due Friday, December 5 at 11:59 PM.”" in r.notes_md
    folder = save(r, tmp_path / "Notes", __import__("datetime").datetime(2026, 10, 5, 13, 0))
    assert folder.name == "2026-10-05 1300 Syllabus"
    assert "*Document · CSE 573 · Syllabus.txt" in (folder / "notes.md").read_text(encoding="utf-8")
    assert "Semantic web basics" in (folder / TEXT_FILE).read_text(encoding="utf-8")
    assert save(r, tmp_path / "Notes", __import__("datetime").datetime(2026, 10, 5, 13, 0)).name.endswith("(2)")
    (tmp_path / "empty.txt").write_text("   ", encoding="utf-8")
    with pytest.raises(ValueError):
        write_notes(tmp_path / "empty.txt", fake)


def test_summarize_course_files_tool_writes_notes_in_the_background(tmp_path):
    from max_assistant.course import CourseLibrary
    from max_assistant.docnotes import DocNotes
    from max_assistant.tools import course as course_tools

    make_course_files(tmp_path / "Course")
    (tmp_path / "Course" / "Week 1.txt").write_text("Introduction to the course and the grading policy for labs.",
                                                    encoding="utf-8")
    reg, ctx = memory_registry()
    notes_dir = tmp_path / "Max Notes"
    lib = CourseLibrary([tmp_path / "Course", notes_dir], WordEmbedder(), tmp_path / "course.db")
    events, finished = [], threading.Event()
    ctx.course = lib
    ctx.docnotes = DocNotes(ctx, lib, notes_dir, FakeSummarizer(), on_event=lambda k, d: events.append((k, d)),
                            on_done=lambda done, failed: finished.set())
    course_tools.register(reg)

    out = reg.run("summarize_course_files", {"file": "pharmacology"})
    assert out.startswith("No course file matches that. The files I have are:") and "Lab 3 handout" in out
    out = reg.run("summarize_course_files", {"course": "CSE 573"})
    assert out.startswith("Writing notes for Lab 3 handout and Lecture 5.")
    assert finished.wait(10)
    notes = ctx.memory.notes()
    assert sorted(n["title"] for n in notes) == ["Lab 3 handout", "Lecture 5"] and {n["kind"] for n in notes} == {"document"}
    assert [d["action"] for k, d in events if k == "docnotes"][-1] == "done"
    # Saved notes are never summarized themselves, and finished files aren't redone
    assert all(notes_dir not in p.parents for p, _ in ctx.docnotes.files())
    assert reg.run("summarize_course_files", {"course": "CSE 573"}).startswith("I already wrote notes for")
    finished.clear()
    out = reg.run("summarize_course_files", {"file": "every file"})
    assert "Midterm review" in out and "Week 1" in out and "already had notes" in out
    assert finished.wait(10) and len(ctx.memory.notes()) == 4
    lib.scan()                                                     # the extracted text isn't indexed twice
    assert not any(p.endswith("document-text.md") for p, *_ in lib._rows)


def test_summarize_everything_in_the_course_folder_goes_straight_to_the_tool():
    from max_assistant.agent import direct_route

    tools = {"summarize_course_files": None}
    assert direct_route("can you summarise each and every file you have in the max course folder and make notes "
                        "of them and put that in the notes section pls", tools) == ("summarize_course_files", {"course": ""})
    assert direct_route("summarize all the cse573 slides", tools) == ("summarize_course_files", {"course": "CSE 573"})
    assert direct_route("summarize this web page", tools) is None
    assert direct_route("what did the professor say about all of it", tools) is None


def test_one_combined_summary_of_a_courses_documents(tmp_path):
    from max_assistant.course import CourseLibrary
    from max_assistant.docnotes import DocNotes, body_of
    from max_assistant.main import documents_done_text
    from max_assistant.tools import course as course_tools

    course = tmp_path / "Course" / "CSE 579"
    course.mkdir(parents=True)
    (course / "Week 1.txt").write_text("Knowledge representation and logic programs. Homework 1 is due "
                                       "September 3 at 11:59 PM.", encoding="utf-8")
    (course / "Week 2.txt").write_text("Answer set programming with clingo and stable models.", encoding="utf-8")
    (tmp_path / "Course" / "Other.txt").write_text("A file from another course entirely.", encoding="utf-8")
    reg, ctx = memory_registry()
    notes_dir = tmp_path / "Max Notes"
    lib = CourseLibrary([tmp_path / "Course", notes_dir], WordEmbedder(), tmp_path / "course.db")
    fake, finished = FakeSummarizer(), threading.Event()
    ctx.course = lib
    ctx.docnotes = DocNotes(ctx, lib, notes_dir, fake, on_done=lambda d, f: finished.set())
    course_tools.register(reg)

    # Week 1 already has notes: they're reused, only Week 2 is written before combining
    assert reg.run("summarize_course_files", {"file": "week 1"}).startswith("Writing notes for Week 1")
    assert finished.wait(10)
    finished.clear()
    out = reg.run("summarize_course_files", {"course": "CSE 579", "combine": True})
    assert out.startswith("Writing notes for Week 2 first, then one combined summary of all 2 files, "
                          "called CSE 579 combined summary.")
    assert finished.wait(10)
    notes = {n["title"]: n for n in ctx.memory.notes()}
    assert sorted(notes) == ["CSE 579 combined summary", "Week 1", "Week 2"]
    combined = notes["CSE 579 combined summary"]
    text = (Path(combined["folder"]) / "notes.md").read_text(encoding="utf-8")
    assert "*Combined summary · CSE 579 · 2 documents" in text and "*Files: Week 1, Week 2*" in text
    assert "- [Week 1, section 1] “Homework 1 is due September 3 at 11:59 PM.”" in text   # quoted, with its file
    assert "### Week 1" in fake.calls[-1][1] and "### Week 2" in fake.calls[-1][1]        # built from both files' notes
    assert "exact quotes" not in fake.calls[-1][1] and "Other" not in text
    assert ctx.docnotes.combined == ["CSE 579 combined summary"]
    assert documents_done_text(["Week 2"], [], ["CSE 579 combined summary"]) == \
        "Your notes for Week 2 are ready in the Notes tab. The CSE 579 combined summary is ready too."
    assert documents_done_text([], [], ["CSE 579 combined summary"]) == \
        "Your CSE 579 combined summary is ready in the Notes tab."
    # Asked again: nothing to redo, so straight to the combined summary
    finished.clear()
    assert reg.run("summarize_course_files", {"course": "CSE 579", "combine": True}).startswith(
        "Writing one combined summary of all 2 files")
    assert finished.wait(10) and len(ctx.memory.notes()) == 4
    assert body_of("# T\n\n*Document · x*\n\n## Summary\nS\n\n## Deadlines & dates (exact quotes)\n- q") == "## Summary\nS"


def test_combined_summary_requests_go_straight_to_the_tool():
    from max_assistant.agent import direct_route

    tools = {"summarize_course_files": None}
    want = ("summarize_course_files", {"course": "CSE 579", "combine": True})
    assert direct_route("summarise all the documents from cse579 at once and create a single summary", tools) == want
    assert direct_route("make one summary of all my CSE 579 slides", tools) == want
    assert direct_route("summarize all the cse573 slides", tools) == ("summarize_course_files", {"course": "CSE 573"})
    assert direct_route("give me one summary of this web page", tools) is None


def test_api_adds_a_document_and_writes_its_notes(tmp_path):
    from max_assistant.course import CourseLibrary
    from max_assistant.docnotes import DocNotes

    client, rt = api_client()
    ctx = rt.ctx
    ctx.course = lib = CourseLibrary([tmp_path / "Course", tmp_path / "Notes"], WordEmbedder(), tmp_path / "course.db")
    finished = threading.Event()
    ctx.docnotes = DocNotes(ctx, lib, tmp_path / "Notes", FakeSummarizer(), on_done=lambda d, f: finished.set())
    assert client.get("/api/documents").json()["files"] == []
    assert client.post("/api/documents/upload?name=notes.exe", content=b"x").status_code == 400
    r = client.post("/api/documents/upload?name=Week 2.txt&course=CSE 573",
                    content=b"Ontologies describe classes and properties. Homework 2 is due October 9 at 5 PM.").json()
    assert r["name"] == "Week 2.txt" and r["summarizing"] and Path(r["saved"]).parent.name == "CSE 573"
    assert finished.wait(10)
    files = client.get("/api/documents").json()["files"]
    assert [(f["name"], f["course"], f["done"]) for f in files] == [("Week 2.txt", "CSE 573", True)]
    note = client.get(f"/api/notes/{ctx.memory.notes()[0]['id']}").json()
    assert note["kind"] == "document" and "Ontologies" in note["transcript_md"] and "Homework 2" in note["notes_md"]
    assert client.post("/api/documents/summarize", json={"paths": [files[0]["path"]]}).json()["skipped"] == ["Week 2.txt"]
    assert client.post("/api/documents/summarize", json={"combine": True}).json()["combined"] == ""   # one file: nothing to combine


# ---------- change Max from the phone (devloop.py) ----------

def test_dev_requests_one_at_a_time_and_need_a_decision(tmp_path, monkeypatch):
    from max_assistant import devloop

    monkeypatch.setattr(devloop, "JOBS_FILE", tmp_path / "jobs.json")
    monkeypatch.setattr(devloop, "DATA", tmp_path)
    monkeypatch.setattr(devloop.DevLoop, "_work", lambda self, job: None)     # no git, no Claude
    loop = devloop.DevLoop()
    with pytest.raises(ValueError):
        loop.submit("hi")
    job = loop.submit("make the orb purple")
    assert job.branch.startswith("phone/1-") and loop.busy is job
    with pytest.raises(RuntimeError, match="still working"):
        loop.submit("another change")
    job.status = "ready"
    with pytest.raises(RuntimeError, match="Approve or reject"):
        loop.submit("another change")
    job.status = "working"
    assert devloop.DevLoop().get(1).status == "failed"            # Max restarted mid-job: never left hanging
    assert devloop.describe_tool("Edit", {"file_path": "C:/x/max_assistant/agent.py"}) == "Editing agent.py"
    assert devloop.describe_tool("Bash", {"command": ".venv/Scripts/python -m pytest -q"}) == "Running the tests"


def test_dev_approve_checks_laptop_edits_and_keeps_the_error(tmp_path, monkeypatch):
    from max_assistant import devloop

    # Raw porcelain: the first line starts with a space, which used to be stripped (cutting a letter off)
    status = " M max_assistant/server.py\nM  config.yaml\nR  old.py -> max_assistant/new.py\n?? tests/new_test.py\n?? data2/\n"
    files = ["max_assistant/server.py", "max_assistant/agent.py", "tests/new_test.py", "max_assistant/new.py",
             "data2/x.txt", "config.yaml"]
    assert devloop.clashes(status, files) == ["config.yaml", "data2/x.txt", "max_assistant/new.py",
                                              "max_assistant/server.py", "tests/new_test.py"]
    assert devloop.clashes("", files) == []

    monkeypatch.setattr(devloop, "JOBS_FILE", tmp_path / "jobs.json")
    monkeypatch.setattr(devloop, "DATA", tmp_path)
    monkeypatch.setattr(devloop.DevLoop, "_work", lambda self, job: None)
    monkeypatch.setattr(devloop, "run", lambda cmd, cwd, timeout=900, env=None: (0, status))
    loop = devloop.DevLoop()
    job = loop.submit("fix the approve button")
    job.status, job.files = "ready", ["max_assistant/server.py"]
    with pytest.raises(RuntimeError, match="unsaved edits on the laptop: max_assistant/server.py"):
        loop.approve(job.id)
    saved = devloop.DevLoop().get(job.id)                          # the phone sees why, after a reload too
    assert saved.status == "ready" and saved.error.startswith("Approve failed: These files have unsaved edits")


def test_dev_api_and_app_update(tmp_path, monkeypatch):
    from max_assistant import devloop

    client, rt = api_client()
    assert client.get("/api/dev/jobs").status_code == 503                 # off unless dev.enabled
    monkeypatch.setattr(devloop, "JOBS_FILE", tmp_path / "jobs.json")
    monkeypatch.setattr(devloop, "DATA", tmp_path)
    monkeypatch.setattr(devloop, "APK", tmp_path / "app" / "max.apk")
    monkeypatch.setattr(devloop.DevLoop, "_work", lambda self, job: None)
    rt.ctx.dev = devloop.DevLoop()
    assert client.post("/api/dev/request", json={"text": "add a dark mode toggle"}).json()["id"] == 1
    assert client.post("/api/dev/request", json={"text": "another"}).status_code == 409
    assert client.get("/api/dev/jobs").json()[0]["request"] == "add a dark mode toggle"
    assert client.post("/api/dev/jobs/1/approve").status_code == 409          # not finished yet
    assert client.post("/api/dev/jobs/1/explode").status_code == 404
    assert client.post("/api/dev/jobs/1/reject").json()["status"] == "rejected"
    assert client.get("/api/app/info").json() == {"available": False}
    assert client.get("/api/app/apk").status_code == 404
    (tmp_path / "app").mkdir()
    (tmp_path / "app" / "max.apk").write_bytes(b"PK fake apk")
    assert client.get("/api/app/info").json()["size"] == 11
    assert client.get("/api/app/apk").content == b"PK fake apk"


# ---------- avatar (max_assistant/avatar, tts envelope, events) ----------

def test_speech_loudness_envelope_follows_the_voice():
    from max_assistant.tts import loudness_envelope

    sr = 16000
    t = np.arange(sr) / sr
    voice = np.concatenate([0.5 * np.sin(2 * np.pi * 220 * t), np.zeros(sr // 2)]).astype(np.float32)
    env = loudness_envelope(voice, sr)
    assert len(env) == 90                                      # 1.5 s at 60 values a second
    assert min(env[:59]) > 0.9 and max(env[61:]) == 0          # loud while talking, shut in the pause
    assert loudness_envelope(np.zeros(10, np.float32), sr) == []


def test_speaker_tells_the_avatar_before_and_after_each_sentence(monkeypatch):
    from max_assistant import audio
    from max_assistant.tts import Speaker

    class Voice:
        def synthesize(self, text):
            return np.full(16000, 0.3, np.float32), 16000

    played, told = [], []
    monkeypatch.setattr(audio, "play", lambda a, sr, device=None: played.append((len(told), len(a))))
    s = Speaker.__new__(Speaker)
    s.piper, s.sapi, s.mic, s.device = Voice(), None, None, None
    s.on_speech = lambda action, data: told.append((action, data))
    s.say("Hello there.")
    assert [a for a, _ in told] == ["start", "end"]
    assert played == [(1, 16000)]                              # the envelope went out before the audio
    assert told[0][1]["rate"] == 60 and len(told[0][1]["env"]) == 60 and told[0][1]["duration"] == 1.0
    assert len(told[0][1]["vis"]) == 60 and len(told[0][1]["vis"][0]) == 5     # mouth shapes per frame
    assert told[0][1]["text"] == "Hello there."                # the avatar's face and hands follow the words
    s.on_speech = lambda action, data: 1 / 0                   # a broken listener never stops speech
    s.say("Still talking.")


def test_mouth_shapes_follow_the_vowel():
    from max_assistant.tts import VISEMES, viseme_track

    sr = 16000
    t = np.arange(sr // 2) / sr

    def vowel(f1, f2):        # a buzz at 120 Hz with two formant-ish peaks
        return sum(np.sin(2 * np.pi * f * t) * a for f, a in ((120, 0.3), (f1, 0.5), (f2, 0.4))).astype(np.float32)

    audio = np.concatenate([vowel(800, 1200), vowel(300, 2400), vowel(350, 700), np.zeros(sr // 4, np.float32)])
    track = viseme_track(audio, sr)
    assert len(track) == len(audio) // (sr // 60) and all(len(f) == 5 for f in track)

    def main_shape(a, b):
        avg = np.mean(track[a:b], axis=0)
        return VISEMES[int(np.argmax(avg))]

    assert main_shape(5, 25) == "aa"                    # open: high F1
    assert main_shape(35, 55) in ("ee", "ih")           # front: high F2
    assert main_shape(65, 85) in ("ou", "oh")           # rounded: low F2
    assert max(max(f) for f in track[-10:]) == 0        # silence: mouth shut
    assert viseme_track(np.zeros(0, np.float32), sr) == []


def test_avatar_remembers_where_it_was_dragged(tmp_path):
    import json

    from max_assistant.avatar import PositionKeeper, saved_position

    path = tmp_path / "pos.json"
    keeper = PositionKeeper(path)
    keeper.check((100, 100, 440, 520))                  # where it started: nothing saved
    keeper.check((100, 100, 440, 520))
    assert not path.exists()
    keeper.check((300, 200, 640, 620))                  # being dragged...
    assert not path.exists()
    keeper.check((300, 200, 640, 620))                  # ...and settled
    assert json.loads(path.read_text()) == {"x": 300, "y": 200, "w": 340, "h": 420}
    keeper.check(None)                                  # hidden: no change
    assert saved_position(path) == (300, 200)
    assert saved_position(path, size=(400, 420)) == (270, 200)   # a wider window: he stays in place
    path.write_text(json.dumps({"x": -99999, "y": -99999}))
    assert saved_position(path) is None                 # that screen is gone
    path.write_text("not json")
    assert saved_position(path) is None


def test_models_asleep_is_published_once_per_change():
    from max_assistant.events import EventBus

    ctx = Context(load_config(), FakeLLM([]))
    ctx.bus = bus = EventBus()
    ctx.models_asleep = True
    ctx.models_asleep = True
    ctx.models_asleep = False
    assert [e["data"] for e in bus.recent if e["kind"] == "gpu"] == [{"asleep": True}, {"asleep": False}]
    assert bus.state["gpu"] == {"asleep": False}               # sent to the avatar when it connects


def test_avatar_hides_for_fullscreen_games_only():
    from max_assistant.avatar import should_hide

    assert should_hide(True, False, "")                        # exclusive fullscreen Direct3D
    assert should_hide(False, True, r"C:\Riot Games\VALORANT\live\VALORANT.exe")
    assert should_hide(False, True, r"D:\SteamLibrary\steamapps\common\Elden Ring\eldenring.exe")
    assert not should_hide(False, True, r"C:\Program Files\Google\Chrome\Application\chrome.exe")   # maximised / video
    assert not should_hide(False, True, r"C:\Windows\explorer.exe")
    assert not should_hide(False, False, r"C:\Riot Games\VALORANT\live\VALORANT.exe")              # windowed
    assert not should_hide(False, True, r"C:\Tools\SomeApp.exe")
    assert should_hide(False, True, r"C:\Tools\MyGame.exe", extra_games=["mygame.exe"])


def test_avatar_outline_runs_become_window_rectangles():
    from max_assistant.avatar import outline_rects

    # 10x10 grid on a 100x200 window: rows 2-4 share a span (one block), row 5 is wider, row 9 has two spans
    runs = [[2, 3, 6], [3, 3, 6], [4, 3, 6], [5, 1, 8], [9, 0, 2], [9, 5, 10]]
    rects = sorted(outline_rects(10, 10, runs, 100, 200))
    assert rects == sorted([(30, 40, 60, 100), (10, 100, 80, 120), (0, 180, 20, 200), (50, 180, 100, 200)])
    assert outline_rects(10, 10, [], 100, 200) == []


# ---------- the phone on its own: reminders / events sync, recordings uploaded later ----------

def test_sync_merge_rule():
    from max_assistant.sync import wins

    assert wins({"status": "done", "updated_ms": 1}, {"status": "pending", "updated_ms": 9})       # finished beats open
    assert not wins({"status": "pending", "updated_ms": 9}, {"status": "cancelled", "updated_ms": 1})
    assert wins({"status": "pending", "updated_ms": 5}, {"status": "pending", "updated_ms": 4})     # later edit wins
    assert not wins({"status": "pending", "updated_ms": 4}, {"status": "pending", "updated_ms": 4})  # tie keeps ours


def test_reminders_get_uids_on_an_old_database():
    import sqlite3

    from max_assistant.memory import MemoryStore
    from max_assistant.reminders import Reminders

    store = MemoryStore(":memory:", WordEmbedder())
    store.db.execute("DROP TABLE reminders")
    store.db.execute("CREATE TABLE reminders (id INTEGER PRIMARY KEY, text TEXT NOT NULL, due TEXT NOT NULL, "
                     "created TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'pending', fired TEXT)")
    store.db.execute("INSERT INTO reminders (text, due, created) VALUES ('old one', '2030-01-01T09:00:00', 'x')")
    r = Reminders(store)
    old = r.pending()[0]
    assert len(old.uid) == 32 and old.updated > 0 and not old.on_phone
    try:
        store.db.execute("INSERT INTO reminders (text, due, created, uid) VALUES ('dup', 'x', 'x', ?)", (old.uid,))
        raise AssertionError("uids must be unique")
    except sqlite3.IntegrityError:
        pass


def test_phone_and_laptop_reminders_sync_both_ways():
    import datetime as dt

    client, rt = api_client()
    laptop = rt.ctx.reminders
    soon = dt.datetime.now().replace(microsecond=0) + dt.timedelta(hours=2)
    ms = lambda d: int(d.timestamp() * 1000)
    mine = laptop.add("set by voice", soon)
    # made on the phone while the laptop was off
    out = client.post("/api/sync", json={"reminders": [
        {"uid": "p1", "text": "buy milk", "due_ms": ms(soon), "status": "pending", "updated_ms": 1000},
        {"uid": mine.uid, "text": "set by voice", "due_ms": ms(mine.due), "status": "done", "updated_ms": 1},
    ]}).json()
    assert laptop.by_uid("p1").text == "buy milk" and laptop.by_uid("p1").due == soon
    assert laptop.by_uid(mine.uid).status == "done"            # the phone fired it: done beats pending
    got = {r["uid"]: r for r in out["reminders"]}
    assert set(got) == {"p1", mine.uid} and got[mine.uid]["status"] == "done"
    assert all(r.on_phone for r in laptop.for_sync())          # the phone holds them all now
    # an older edit from the phone doesn't undo a newer one on the laptop
    laptop.set_status(laptop.by_uid("p1").id, "cancelled")
    client.post("/api/sync", json={"reminders": [
        {"uid": "p1", "text": "buy oat milk", "due_ms": ms(soon), "status": "pending", "updated_ms": 2000}]})
    assert laptop.by_uid("p1").status == "cancelled"
    assert client.post("/api/sync", json={"reminders": [{"uid": "", "text": "x", "due_ms": 0}]}).status_code == 200


def test_missed_reminders_the_phone_showed_stay_quiet():
    import datetime as dt

    from max_assistant.reminders import ReminderScheduler, Reminders

    r = Reminders(memory_store())
    now = dt.datetime.now()
    shown = r.add("phone had this", now - dt.timedelta(hours=3))
    r.add("laptop only", now - dt.timedelta(hours=3))
    on_time = r.add("on time", now - dt.timedelta(seconds=20))
    r.mark_on_phone([shown.uid, on_time.uid])
    fired = []
    ReminderScheduler(r, lambda rem, missed: fired.append((rem.text, missed))).check(now)
    assert sorted(fired) == [("laptop only", True), ("on time", False)]   # due right now: both devices ring
    assert r.pending() == []


def test_events_tools_api_and_sync():
    import datetime as dt

    client, rt = api_client()
    from max_assistant.tools import memory as memory_tools

    reg = ToolRegistry(context=rt.ctx)
    memory_tools.register(reg)
    out = reg.run("add_event", {"title": "Study group", "when": "tomorrow at 3pm", "duration_minutes": 90})
    assert out == "Added to your calendar: Study group, tomorrow at 3 PM."
    assert reg.run("add_event", {"title": "Mom's birthday", "when": "in 3 days"}).startswith("Added")
    assert reg.run("add_event", {"title": "Picnic", "when": "next saturday"}).endswith("(all day).")
    assert reg.run("add_event", {"title": "x", "when": "blorp"}).startswith("Error")
    upcoming = rt.ctx.events.upcoming(30)
    study = next(x for x in upcoming if x.title == "Study group")
    assert (study.end - study.start) == dt.timedelta(minutes=90) and not study.all_day
    picnic = next(x for x in upcoming if x.title == "Picnic")
    assert picnic.all_day and picnic.start.hour == 0 and picnic.end - picnic.start == dt.timedelta(days=1)
    assert reg.run("list_events", {"day": "tomorrow"}).startswith("On your calendar tomorrow: Study group")
    assert "Study group" in [x["title"] for x in client.get("/api/events").json()]
    # an event made on the phone offline, and the phone deleting one of the laptop's
    start = dt.datetime.now().replace(microsecond=0) + dt.timedelta(days=2)
    ms = lambda d: int(d.timestamp() * 1000)
    out = client.post("/api/sync", json={"events": [
        {"uid": "pe1", "title": "Dentist", "start_ms": ms(start), "end_ms": ms(start) + 1800_000, "all_day": False,
         "location": "Tempe", "status": "active", "updated_ms": ms(dt.datetime.now())},
        {**study.to_sync(), "status": "cancelled", "updated_ms": study.updated + 1}]}).json()
    assert rt.ctx.events.by_uid("pe1").location == "Tempe"
    assert rt.ctx.events.by_uid(study.uid).status == "cancelled"
    assert {x["uid"] for x in out["events"]} >= {"pe1", study.uid, picnic.uid}   # cancelled sent: the phone drops it
    assert reg.run("cancel_event", {"which": "dentist"}).startswith("Removed from your calendar: Dentist")
    assert "events" in client.get("/api/today").json()


def test_phone_recording_upload_becomes_notes(tmp_path):
    import datetime as dt

    import numpy as np

    from max_assistant.notes import NotesManager
    from max_assistant.phone_notes import PhoneRecordings

    client, rt = api_client()
    ctx = rt.ctx
    heard = []

    def transcriber():
        def transcribe(audio, prompt):
            heard.append(len(audio))
            return "The homework is due Friday at 11:59 PM. Today we cover sorting."
        return transcribe

    ctx.notes = NotesManager(ctx, transcriber, tmp_path / "notes")
    ctx.notes.summarize = lambda system, text, n: "## Summary\nSorting lecture."
    events = []
    ctx.phone_notes = pn = PhoneRecordings(ctx, tmp_path / "rec", on_event=lambda k, d: events.append((k, d)),
                                           decoder=lambda path: np.full(16000 * 70, 0.1, np.float32))
    pn.start = lambda: pn                     # no background thread: process() is called below
    started = dt.datetime(2026, 10, 9, 9, 30)
    r = client.post("/api/notes/upload?uid=rec-1&kind=lecture&started_ms=" + str(int(started.timestamp() * 1000)),
                    content=b"fake aac bytes")
    assert r.status_code == 200 and r.json()["state"] == "queued" and r.json()["title"] == "Lecture Oct 9, 9:30 AM"
    assert pn.pending() == ["rec-1"]
    st = pn.process("rec-1")
    assert st["state"] == "done" and heard and sum(heard) == 16000 * 70
    note = ctx.memory.note(st["note_id"])
    assert note["started"] == "2026-10-09T09:30:00"           # dated when it was recorded
    assert "Sorting lecture" in note["summary"]
    assert pn.pending() == [] and not list((tmp_path / "rec").glob("*.audio"))   # no audio kept
    assert client.get("/api/notes/upload/rec-1").json()["note_id"] == st["note_id"]
    assert ("notes", "saved") in [(k, d["action"]) for k, d in events]
    # uploading again (the phone didn't hear back) changes nothing
    again = client.post("/api/notes/upload?uid=rec-1", content=b"fake aac bytes").json()
    assert again["state"] == "done" and pn.pending() == []
    assert client.get("/api/notes/upload/nope").status_code == 404
    assert client.post("/api/notes/upload?uid=..%2F", content=b"x").status_code == 400
    # a broken file: kept (renamed) and reported; a new upload can retry
    def broken(path):
        raise RuntimeError("bad audio")
    pn.decoder = broken
    client.post("/api/notes/upload?uid=rec-2", content=b"junk")
    assert pn.process("rec-2")["state"] == "failed" and (tmp_path / "rec" / "rec-2.failed").exists()
    assert client.post("/api/notes/upload?uid=rec-2", content=b"junk").json()["state"] == "queued"


def test_decoding_an_aac_recording(tmp_path):
    """The phone records AAC (ADTS); PyAV (faster-whisper's decoder) reads it."""
    import av
    import numpy as np

    from max_assistant.phone_notes import decode

    path = tmp_path / "x.aac"
    with av.open(str(path), "w", format="adts") as out:
        stream = out.add_stream("aac", rate=16000)
        stream.layout = "mono"
        t = np.arange(16000 * 2) / 16000
        pcm = (np.sin(2 * np.pi * 440 * t) * 0.3).astype(np.float32)
        for i in range(0, len(pcm), 1024):
            frame = av.AudioFrame.from_ndarray(pcm[None, i:i + 1024], format="flt", layout="mono")
            frame.sample_rate = 16000
            for packet in stream.encode(frame):
                out.mux(packet)
        for packet in stream.encode(None):
            out.mux(packet)
    audio = decode(path)
    assert abs(len(audio) / 16000 - 2.0) < 0.2 and float(np.abs(audio).max()) > 0.1
