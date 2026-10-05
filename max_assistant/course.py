"""Ask your course material: slides, PDFs, readings and handouts, plus saved lecture notes.

Files in the course folder (one subfolder per course is optional) are split into passages by
page or slide, embedded with the same small CPU model as memory (bge-small) and kept in
data/course.db. A background scan picks up new or changed files. Questions retrieve the best
passages (meaning plus a small boost for exact terms like "SPARQL") and the model answers
from them, naming the file and page. The files themselves are never changed or moved.
"""
from __future__ import annotations

import logging
import re
import sqlite3
import threading
import time
from dataclasses import dataclass
from pathlib import Path

import numpy as np

from .config import resolve_path

log = logging.getLogger(__name__)
TYPES = {".pdf", ".pptx", ".docx", ".txt", ".md"}
SKIP_NAMES = {"document-text.md"}          # text Max extracted when writing a document's notes (docnotes.TEXT_FILE)
SCHEMA = """
CREATE TABLE IF NOT EXISTS files (path TEXT PRIMARY KEY, mtime REAL NOT NULL, size INTEGER NOT NULL, course TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS chunks (
    id INTEGER PRIMARY KEY, path TEXT NOT NULL, course TEXT NOT NULL, page TEXT NOT NULL, text TEXT NOT NULL, vec BLOB NOT NULL);
"""


@dataclass
class Passage:
    file: str          # file name
    course: str
    page: str          # "p. 4", "slide 12", "section 3"
    text: str
    score: float


def pages(path: Path) -> list[tuple[str, str]]:
    """(location, text) per page / slide / section."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        import pymupdf

        with pymupdf.open(path) as doc:
            return [(f"p. {i + 1}", page.get_text()) for i, page in enumerate(doc)]
    if ext == ".pptx":
        from pptx import Presentation

        out = []
        for i, slide in enumerate(Presentation(str(path)).slides):
            texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text]
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text:
                texts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
            out.append((f"slide {i + 1}", "\n".join(texts)))
        return out
    if ext == ".docx":
        import docx

        paras = [p.text for p in docx.Document(str(path)).paragraphs if p.text.strip()]
        return [(f"section {i // 12 + 1}", "\n".join(paras[i:i + 12])) for i in range(0, len(paras), 12)]
    text = path.read_text(encoding="utf-8", errors="ignore")
    blocks = re.split(r"\n(?=#{1,3} )", text)              # Markdown headings, e.g. lecture notes
    return [(f"section {i + 1}", b) for i, b in enumerate(blocks)]


def chunk(text: str, words: int = 160, overlap: int = 30) -> list[str]:
    w = text.split()
    if len(w) <= words:
        return [" ".join(w)] if len(w) >= 8 else []
    step = words - overlap
    return [" ".join(w[i:i + words]) for i in range(0, max(len(w) - overlap, 1), step) if len(w[i:i + words]) >= 8]


class CourseLibrary:
    def __init__(self, folders: list[Path], embedder, db_path: str | Path = "data/course.db"):
        self.folders = [Path(f) for f in folders]
        self.embedder = embedder
        self.index_path = resolve_path(db_path)
        self.index_path.parent.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(str(self.index_path), check_same_thread=False)
        self.db.executescript(SCHEMA)
        self.lock = threading.RLock()
        self._load()
        self._stop = threading.Event()
        self.scanning = False

    def _load(self):
        with self.lock:
            rows = self.db.execute("SELECT id, path, course, page, text, vec FROM chunks").fetchall()
            self._rows = [(r[1], r[2], r[3], r[4]) for r in rows]
            self._vecs = np.stack([np.frombuffer(r[5], dtype=np.float32) for r in rows]) if rows else np.zeros((0, 384), np.float32)

    def _course_of(self, path: Path, root: Path) -> str:
        rel = path.relative_to(root)
        return rel.parts[0] if len(rel.parts) > 1 else root.name

    def files(self) -> list[tuple[Path, str]]:
        """(path, course) of every readable file in the folders, sorted by course then name."""
        out = []
        for root in self.folders:
            if not root.is_dir():
                continue
            for f in root.rglob("*"):
                if f.is_file() and f.suffix.lower() in TYPES and not f.name.startswith(("~$", ".")) \
                        and f.name not in SKIP_NAMES:
                    out.append((f, self._course_of(f, root)))
        return sorted(out, key=lambda pc: (pc[1].lower(), pc[0].name.lower()))

    def scan(self) -> dict:
        """Index new or changed files, drop deleted ones. Returns counts."""
        self.scanning = True
        added = removed = 0
        try:
            seen: set[str] = set()
            for f, course in self.files():
                key = str(f.resolve())
                seen.add(key)
                st = f.stat()
                with self.lock:
                    row = self.db.execute("SELECT mtime, size FROM files WHERE path = ?", (key,)).fetchone()
                if row and abs(row[0] - st.st_mtime) < 1 and row[1] == st.st_size:
                    continue
                try:
                    self._index_file(f, key, course, st)
                    added += 1
                except Exception as exc:
                    log.warning("couldn't read %s: %s", f.name, exc)
            with self.lock:
                gone = [r[0] for r in self.db.execute("SELECT path FROM files") if r[0] not in seen]
                for key in gone:
                    self.db.execute("DELETE FROM chunks WHERE path = ?", (key,))
                    self.db.execute("DELETE FROM files WHERE path = ?", (key,))
                    removed += 1
                self.db.commit()
            if added or removed:
                self._load()
                log.info("course material: %d files indexed, %d removed, %d passages", added, removed, len(self._rows))
        finally:
            self.scanning = False
        return {"indexed": added, "removed": removed, "passages": len(self._rows)}

    def _index_file(self, f: Path, key: str, course: str, st):
        items = [(loc, piece) for loc, text in pages(f) for piece in chunk(text)]
        vecs = self.embedder.embed([f"{f.stem} ({loc}): {t}" for loc, t in items]) if items else []
        with self.lock:
            self.db.execute("DELETE FROM chunks WHERE path = ?", (key,))
            self.db.executemany("INSERT INTO chunks (path, course, page, text, vec) VALUES (?, ?, ?, ?, ?)",
                                [(key, course, loc, t, np.asarray(v, np.float32).tobytes()) for (loc, t), v in zip(items, vecs)])
            self.db.execute("INSERT OR REPLACE INTO files (path, mtime, size, course) VALUES (?, ?, ?, ?)",
                            (key, st.st_mtime, st.st_size, course))
            self.db.commit()

    def search(self, query: str, k: int = 5, course: str = "", min_score: float = 0.45) -> list[Passage]:
        with self.lock:
            if not self._rows:
                return []
            q = self.embedder.embed([query])[0]
            scores = self._vecs @ q
            terms = [t for t in re.findall(r"[a-z0-9]{3,}", query.lower()) if t not in STOP]
            for i, (_, c, _, text) in enumerate(self._rows):
                low = text.lower()
                scores[i] += 0.04 * sum(1 for t in terms if t in low)       # exact terms (SPARQL, RDF...) count
                if course and course.lower().replace(" ", "") not in c.lower().replace(" ", ""):
                    scores[i] -= 1.0
            best = np.argsort(-scores)[: k * 3]
            out, seen = [], set()
            for i in best:
                path, c, page, text = self._rows[i]
                if scores[i] < min_score or (path, page) in seen:
                    continue
                seen.add((path, page))
                out.append(Passage(Path(path).name, c, page, text, float(scores[i])))
                if len(out) == k:
                    break
            return out

    def stats(self) -> dict:
        with self.lock:
            files = self.db.execute("SELECT course, COUNT(*) FROM files GROUP BY course").fetchall()
        return {"courses": {c: n for c, n in files}, "passages": len(self._rows), "folders": [str(f) for f in self.folders]}

    def clear(self):
        with self.lock:
            self.db.execute("DELETE FROM chunks")
            self.db.execute("DELETE FROM files")
            self.db.commit()
            self.db.execute("VACUUM")
        self._load()

    def start(self, every_s: float = 600):
        def run():
            while True:
                try:
                    self.scan()
                except Exception as exc:
                    log.warning("course material scan failed: %s", exc)
                if self._stop.wait(every_s):
                    return
        threading.Thread(target=run, name="course-scan", daemon=True).start()
        return self

    def stop(self):
        self._stop.set()


STOP = {"the", "and", "what", "did", "say", "about", "from", "for", "with", "this", "that", "how", "does", "are", "was",
        "professor", "lecture", "slides", "explain", "tell", "according", "notes", "class", "course", "handout", "reading"}
